from pathlib import Path
from shutil import copy2
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
PPTX_PATH = OUT_DIR / "schule_wohlbefinden_praesentation.pptx"
KEYNOTE_PPTX_PATH = OUT_DIR / "schule_wohlbefinden_keynote_ipad.pptx"
GUIDE_PATH = OUT_DIR / "sprecherleitfaden_schule_wohlbefinden.md"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS = {"p": P_NS, "a": A_NS, "mc": MC_NS, "p14": P14_NS}


COLORS = {
    "cream": RGBColor(255, 250, 241),
    "navy": RGBColor(35, 48, 86),
    "muted": RGBColor(91, 102, 130),
    "blue": RGBColor(190, 229, 245),
    "blue2": RGBColor(220, 244, 252),
    "pink": RGBColor(247, 203, 218),
    "pink2": RGBColor(255, 230, 238),
    "yellow": RGBColor(255, 240, 179),
    "yellow2": RGBColor(255, 248, 216),
    "mint": RGBColor(215, 242, 225),
    "mint2": RGBColor(236, 250, 241),
    "lavender": RGBColor(226, 217, 250),
    "lavender2": RGBColor(244, 239, 255),
    "white": RGBColor(255, 255, 255),
    "line": RGBColor(230, 218, 205),
    "red_soft": RGBColor(250, 174, 174),
}


SLIDES = [
    {
        "title": "Wie Schule wirkt",
        "subtitle": "Strukturelle Unterschiede und ihr Effekt auf das Wohlbefinden von Schülerinnen und Schülern",
        "section": "Seminarfachkurs Glück",
        "time": "0:00-0:45",
        "speaker": [
            "Ich stelle die Leitfrage der Facharbeit vor: Inwiefern beeinflussen strukturelle Unterschiede im deutschen Schulsystem das Wohlbefinden von Schülerinnen und Schülern?",
            "Wichtig ist: Schule wird hier nicht nur als Lernort betrachtet, sondern als Lebensraum, der Selbstwert, Motivation und Zukunftsperspektiven prägt.",
        ],
    },
    {
        "title": "Warum ist das Thema relevant?",
        "section": "Einstieg",
        "time": "0:45-1:45",
        "cards": [
            ("Schule als Lebensraum", "Schule vermittelt Wissen, qualifiziert und integriert - sie ist aber auch ein Ort intensiver Bewertung.", "blue"),
            ("Belastung im Jugendalter", "Leistungsdruck, Prüfungsstress und Vergleichsmechanismen werden in der Arbeit als zentrale Belastungsfaktoren beschrieben.", "pink"),
            ("Bildungsgerechtigkeit", "Der Bildungserfolg hängt in Deutschland vergleichsweise stark mit dem sozioökonomischen Hintergrund zusammen.", "yellow"),
        ],
        "speaker": [
            "Die Arbeit begründet das Thema mit drei Punkten: Schule hat einen Bildungsauftrag, ist aber zugleich Teil der Lebensrealität Jugendlicher.",
            "Wenn Schule stark bewertet und auswählt, kann das nicht nur Bildungswege, sondern auch Wohlbefinden beeinflussen.",
            "Dazu kommt die Frage nach Bildungsgerechtigkeit, weil soziale Herkunft laut der Arbeit eine große Rolle für Bildungschancen spielt.",
        ],
    },
    {
        "title": "Leitfrage und Vorgehen",
        "section": "Fokus",
        "time": "1:45-2:45",
        "quote": "Inwiefern beeinflussen strukturelle Unterschiede im deutschen Schulsystem das Wohlbefinden von Schülerinnen und Schülern?",
        "bullets": [
            "literaturgestützte Analyse",
            "Grundlagen: Schulstruktur, Wohlbefinden, Bildungsgerechtigkeit",
            "Analyse: deutsches System, Leistungsdruck, Schulklima, internationaler Vergleich",
            "Bewertung: Stärken, Schwächen und Reformansätze",
        ],
        "speaker": [
            "Die Leitfrage verbindet also zwei Bereiche: Strukturen des Schulsystems und das Wohlbefinden der Lernenden.",
            "Die Facharbeit nutzt keine eigene Umfrage, sondern wertet Fachliteratur, Bildungsberichte, OECD-Daten und gesundheitswissenschaftliche Studien aus.",
            "Für die Präsentation gehe ich deshalb von den Begriffen über das deutsche System bis zum Fazit.",
        ],
    },
    {
        "title": "Drei Grundbegriffe",
        "section": "Theoretische Grundlage",
        "time": "2:45-3:45",
        "cards": [
            ("Strukturmerkmale", "Schulformen, Dauer gemeinsamer Schulzeit, Noten, Übergänge, Versetzungen und Abschlussprüfungen.", "blue"),
            ("Wohlbefinden", "Psychische und soziale Dimensionen: Lebenszufriedenheit, Stress, Sicherheit, Eingebundenheit und Selbstwirksamkeit.", "pink"),
            ("Bildungsgerechtigkeit", "Faire Bildungschancen unabhängig von sozialer Herkunft, Geschlecht oder Migrationshintergrund.", "yellow"),
        ],
        "speaker": [
            "Damit der Zusammenhang verständlich wird, klärt die Arbeit drei Begriffe.",
            "Strukturmerkmale sind die Rahmenbedingungen, die Schule organisiert.",
            "Wohlbefinden meint nicht nur keine Krankheit, sondern auch emotionale Sicherheit, soziale Eingebundenheit und Selbstwirksamkeit.",
            "Bildungsgerechtigkeit ist wichtig, weil frühe Auswahl und Bewertungen je nach sozialem Hintergrund unterschiedlich wirken können.",
        ],
    },
    {
        "title": "Welche Funktionen erfüllt Schule?",
        "section": "Theoretische Grundlage",
        "time": "3:45-4:45",
        "functions": [
            ("Qualifikation", "Wissen und Kompetenzen für Beruf, Studium und Teilhabe"),
            ("Selektion", "Zuordnung zu Bildungswegen durch Noten, Empfehlungen und Abschlüsse"),
            ("Sozialisation", "Vermittlung von Normen, Werten und Verhaltensweisen"),
            ("Integration & Legitimation", "Zugehörigkeit schaffen und Bildungsabschlüsse gesellschaftlich begründen"),
        ],
        "speaker": [
            "Die Facharbeit zeigt, dass Schule mehrere Funktionen gleichzeitig erfüllt.",
            "Besonders spannend für das Thema Glück und Wohlbefinden ist die Spannung zwischen Qualifikation und Selektion.",
            "Schule soll fördern und vorbereiten, sortiert aber auch über Noten und Abschlüsse. Genau daraus können Druck und Ungleichheiten entstehen.",
        ],
    },
    {
        "title": "Das deutsche Schulsystem: frühe Differenzierung",
        "section": "Struktur",
        "time": "4:45-5:55",
        "flow": True,
        "speaker": [
            "Charakteristisch für Deutschland ist die Mehrgliedrigkeit.",
            "Nach der Grundschule erfolgt in vielen Bundesländern eine Aufteilung in verschiedene Schulformen.",
            "Gymnasium, Realschule, Hauptschule, Gesamtschule und berufliche Schulen verfolgen unterschiedliche Ziele und eröffnen unterschiedliche Wege.",
            "Die Arbeit betont: Diese Differenzierung kann Förderung ermöglichen, legt Bildungswege aber auch früh fest.",
        ],
    },
    {
        "title": "Mechanismen, die Druck erzeugen können",
        "section": "Leistungsdruck",
        "time": "5:55-7:00",
        "mechanisms": [
            ("Noten", "Vergleichbarkeit und Rückmeldung - aber auch Konkurrenzdruck"),
            ("Versetzung", "Leistungen entscheiden über den nächsten Schritt"),
            ("Sitzenbleiben", "Kann Motivation und Selbstwertgefühl negativ beeinflussen"),
            ("soziale Herkunft", "Bildungserfolg ist laut Arbeit eng mit familiären Ressourcen verbunden"),
        ],
        "speaker": [
            "Neben der Schulform prägen konkrete Selektionsmechanismen den Alltag.",
            "Noten erfüllen wichtige Funktionen, können aber auch Vergleichskultur und Konkurrenz verstärken.",
            "Sitzenbleiben wird in der Arbeit als Beispiel genannt, weil es nicht immer bessere Lernergebnisse bringt und Motivation sowie Selbstwert belasten kann.",
            "Außerdem wird deutlich: Bildungschancen sind in Deutschland laut OECD-Daten stark mit sozialer Herkunft verknüpft.",
        ],
    },
    {
        "title": "Kurze Kursaktivierung",
        "section": "Publikum einbinden",
        "time": "7:00-8:30",
        "interaction": True,
        "speaker": [
            "Jetzt kommt eine kurze Einbindung: Ich frage den Kurs, welcher Faktor das eigene Wohlbefinden in Schule am stärksten beeinflusst.",
            "Die vier Antwortmöglichkeiten stammen direkt aus den Themen der Facharbeit: Leistungsdruck, Beziehung zu Lehrkräften, Mitbestimmung und Lebensweltbezug.",
            "Nach einer kurzen Handabstimmung greife ich zwei bis drei Stimmen auf und leite damit zum Kapitel Schule und Wohlbefinden über.",
        ],
    },
    {
        "title": "Schule und Wohlbefinden",
        "section": "Wirkung",
        "time": "8:30-9:45",
        "wellbeing": True,
        "speaker": [
            "Die Arbeit beschreibt Schule als zentralen sozialen Lebensraum.",
            "Leistungsdruck kann mit Stress, Angst vor Versagen und psychosomatischen Beschwerden wie Schlafproblemen, Kopfschmerzen oder Konzentrationsschwierigkeiten verbunden sein.",
            "Gleichzeitig können ein positives Schulklima, unterstützende Lehrkräfte, Mitbestimmung und gute Feedbackkultur Wohlbefinden und Motivation stärken.",
        ],
    },
    {
        "title": "Lebensweltbezug: Warum Inhalte wichtig sind",
        "section": "Unterricht",
        "time": "9:45-10:45",
        "cards": [
            ("Alltagskompetenzen", "In der Arbeit werden Steuererklärungen, Versicherungen, Mietverträge und Finanzplanung als Beispiele genannt.", "yellow"),
            ("Finanzbildung", "OECD-Studien zu Financial Literacy zeigen, dass finanzielle Grundkompetenzen international ein Thema sind.", "blue"),
            ("Berufsorientierung", "Praktika und Orientierung können realistische Vorstellungen von Bildungs- und Berufswegen fördern.", "pink"),
        ],
        "speaker": [
            "Wohlbefinden hängt nicht nur mit Druck zusammen, sondern auch damit, ob Unterricht als sinnvoll erlebt wird.",
            "Die Arbeit nennt Alltagskompetenzen wie Finanzplanung oder Mietverträge als Bereiche, die Jugendliche häufig als lebensnah ansehen.",
            "Auch Berufsorientierung und Praktika können Unsicherheit über die Zukunft reduzieren.",
        ],
    },
    {
        "title": "Internationaler Vergleich",
        "section": "Einordnung",
        "time": "10:45-12:00",
        "comparison": True,
        "speaker": [
            "Der internationale Vergleich hilft, die deutsche Struktur einzuordnen.",
            "Skandinavische Modelle wie Finnland setzen auf längeres gemeinsames Lernen und spätere Differenzierung.",
            "Angelsächsische Systeme arbeiten häufiger mit umfassenden Gesamtschulen, Wahlmöglichkeiten und projektorientiertem Lernen.",
            "Die Arbeit betont aber auch: Es gibt kein perfektes System. Jedes Bildungssystem hat eigene Herausforderungen.",
        ],
    },
    {
        "title": "Stärken und Schwächen des deutschen Systems",
        "section": "Bewertung",
        "time": "12:00-13:15",
        "proscons": True,
        "speaker": [
            "In der Bewertung arbeitet die Facharbeit sowohl Stärken als auch Schwächen heraus.",
            "Stärken sind die Differenzierung verschiedener Bildungswege, das duale Ausbildungssystem und hohe akademische Standards im Gymnasium.",
            "Kritisch bewertet werden frühe Aufteilung, soziale Herkunftseffekte, Leistungsdruck und teilweise soziale Segregation zwischen Schulformen.",
        ],
    },
    {
        "title": "Reformansätze aus der Arbeit",
        "section": "Ausblick",
        "time": "13:15-14:10",
        "reforms": [
            ("Mehr Lebenskompetenzen", "Finanzbildung, Medienkompetenz und praktische Alltagsfähigkeiten stärker integrieren"),
            ("Spätere Differenzierung", "Mehr Zeit für Entwicklung und fundiertere Bildungsentscheidungen geben"),
            ("Alternative Bewertung", "Lernfeedback, Portfolioarbeit und projektbasierte Bewertung stärker nutzen"),
        ],
        "speaker": [
            "Als Reformansätze nennt die Arbeit drei Richtungen.",
            "Erstens mehr Lebenskompetenzen im Curriculum, zweitens eine spätere Leistungsdifferenzierung und drittens alternative Bewertungsformen.",
            "Diese Vorschläge sollen Bildungsqualität und Schülerwohlbefinden stärker zusammenbringen.",
        ],
    },
    {
        "title": "Fazit: Was beantwortet die Leitfrage?",
        "section": "Schluss",
        "time": "14:10-15:00",
        "takeaways": [
            "Schulische Strukturen beeinflussen Lernbedingungen, Bildungswege und subjektives Wohlbefinden.",
            "Frühe Differenzierung und Leistungsbewertung können fördern, aber auch Druck und Ungleichheit verstärken.",
            "Schulklima, Unterstützung, Mitbestimmung und Lebensweltbezug können Wohlbefinden deutlich stärken.",
            "Eine Schule der Zukunft sollte Leistung, Gerechtigkeit und persönliche Entwicklung gemeinsam denken.",
        ],
        "speaker": [
            "Die Antwort auf die Leitfrage lautet: Strukturen wirken deutlich auf das Wohlbefinden, weil sie bestimmen, wie Schülerinnen und Schüler Leistung, Chancen, Beziehungen und Zukunft erleben.",
            "Das deutsche System hat Stärken, aber auch belastende Seiten.",
            "Für den Seminarfachkurs Glück ist besonders wichtig: Wohlbefinden entsteht nicht nur individuell, sondern wird durch schulische Rahmenbedingungen mitgestaltet.",
        ],
    },
]


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color


def add_textbox(slide, x, y, w, h, text, size=24, bold=False, color=None, align=PP_ALIGN.LEFT):
    color = color or COLORS["navy"]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos Display"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_multiline(slide, x, y, w, h, lines, size=20, color=None, bullet=False, spacing=0):
    color = color or COLORS["navy"]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.06)
    tf.margin_bottom = Inches(0.06)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.space_after = Pt(spacing)
        if bullet:
            p.text = f"- {line}"
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
    return box


def add_decor(slide, accent="blue"):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS["cream"]

    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
    set_fill(band, COLORS[accent])

    for x, y, size, color in [
        (11.45, -0.35, 1.6, "pink2"),
        (12.35, 0.55, 0.95, "yellow2"),
        (-0.35, 6.65, 1.25, "blue2"),
    ]:
        oval = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
        set_fill(oval, COLORS[color])


def add_header(slide, title, section=None, accent="blue"):
    add_decor(slide, accent)
    if section:
        pill = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(0.42), Inches(2.3), Inches(0.42))
        set_fill(pill, COLORS[accent])
        add_textbox(slide, 0.85, 0.48, 2.1, 0.22, section.upper(), size=9, bold=True, color=COLORS["navy"], align=PP_ALIGN.CENTER)
    add_textbox(slide, 0.75, 0.82, 11.7, 0.58, title, size=28, bold=True)


def add_footer(slide, index, time):
    add_textbox(slide, 0.75, 7.02, 3.5, 0.22, f"Folie {index:02d} | ca. {time}", size=8, color=COLORS["muted"])
    add_textbox(slide, 10.25, 7.02, 2.25, 0.22, "Seminarfachkurs Glück", size=8, color=COLORS["muted"], align=PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, title, body, fill_key):
    shadow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x + 0.04), Inches(y + 0.06), Inches(w), Inches(h))
    set_fill(shadow, COLORS["line"])
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(card, COLORS[fill_key])
    add_textbox(slide, x + 0.18, y + 0.18, w - 0.36, 0.32, title, size=16, bold=True)
    add_multiline(slide, x + 0.18, y + 0.62, w - 0.36, h - 0.75, [body], size=13, color=COLORS["navy"])


def add_title_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_decor(slide, "blue")
    big = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.82), Inches(0.88), Inches(11.7), Inches(4.85))
    set_fill(big, COLORS["white"])
    circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(8.9), Inches(0.2), Inches(2.8), Inches(2.8))
    set_fill(circle, COLORS["pink"])
    circle2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.95), Inches(1.83), Inches(1.35), Inches(1.35))
    set_fill(circle2, COLORS["yellow"])
    add_textbox(slide, 1.25, 1.18, 7.7, 0.62, data["section"], size=14, bold=True, color=COLORS["muted"])
    add_textbox(slide, 1.25, 1.82, 8.5, 0.98, data["title"], size=44, bold=True)
    add_multiline(slide, 1.28, 3.02, 7.6, 0.92, [data["subtitle"]], size=21, color=COLORS["navy"])
    add_textbox(slide, 1.28, 4.35, 4.6, 0.38, "Facharbeit von Svea Timphus", size=15, color=COLORS["muted"])
    add_textbox(slide, 1.28, 4.82, 4.8, 0.28, "Leitfrage, zentrale Ergebnisse und Ausblick", size=12, color=COLORS["muted"])
    add_footer(slide, 1, data["time"])


def add_cards_slide(prs, idx, data, accent):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, data["title"], data.get("section"), accent)
    cards = data["cards"]
    for pos, card in enumerate(cards):
        add_card(slide, 0.82 + pos * 4.08, 2.05, 3.58, 3.35, *card)
    add_footer(slide, idx, data["time"])


def add_leitfrage_slide(prs, idx, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, data["title"], data.get("section"), "pink")
    quote = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.75), Inches(11.3), Inches(1.45))
    set_fill(quote, COLORS["pink2"])
    add_textbox(slide, 1.35, 2.03, 10.6, 0.78, f"\"{data['quote']}\"", size=22, bold=True, align=PP_ALIGN.CENTER)
    add_multiline(slide, 1.25, 3.75, 10.8, 1.8, data["bullets"], size=18, bullet=True, color=COLORS["navy"], spacing=5)
    add_footer(slide, idx, data["time"])


def add_functions_slide(prs, idx, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, data["title"], data.get("section"), "yellow")
    colors = ["blue", "pink", "yellow", "mint"]
    for i, (title, body) in enumerate(data["functions"]):
        x = 1.1 + (i % 2) * 5.65
        y = 2.0 + (i // 2) * 2.05
        bubble = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(0.75), Inches(0.75))
        set_fill(bubble, COLORS[colors[i]])
        add_textbox(slide, x + 0.05, y + 0.19, 0.65, 0.24, str(i + 1), size=15, bold=True, align=PP_ALIGN.CENTER)
        add_card(slide, x + 0.95, y - 0.08, 4.35, 1.2, title, body, f"{colors[i]}2" if colors[i] != "yellow" else "yellow2")
    add_footer(slide, idx, data["time"])


def add_flow_slide(prs, idx, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, data["title"], data.get("section"), "blue")
    add_textbox(slide, 0.95, 1.62, 3.2, 0.42, "Grundschule", size=19, bold=True, align=PP_ALIGN.CENTER)
    root = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(2.1), Inches(3.2), Inches(0.72))
    set_fill(root, COLORS["blue"])
    add_textbox(slide, 1.1, 2.3, 2.9, 0.22, "gemeinsamer Start", size=14, bold=True, align=PP_ALIGN.CENTER)
    arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(4.32), Inches(2.18), Inches(1.15), Inches(0.52))
    set_fill(arrow, COLORS["lavender"])
    add_textbox(slide, 4.03, 2.92, 1.7, 0.42, "oft nach Klasse 4", size=11, color=COLORS["muted"], align=PP_ALIGN.CENTER)
    x_positions = [5.85, 7.62, 9.39, 11.16]
    labels = [
        ("Gymnasium", "Abitur, stärker theoretisch, hoher Prüfungsdruck möglich", "pink2"),
        ("Realschule", "Mittlerer Abschluss, Verbindung von Allgemeinbildung und Praxis", "yellow2"),
        ("Hauptschule", "stärker praktisch, teils mit Stigmatisierung verbunden", "mint2"),
        ("Gesamtschule", "längeres gemeinsames Lernen, mehrere Abschlüsse", "lavender2"),
    ]
    for x, (title, body, color) in zip(x_positions, labels):
        add_card(slide, x, 1.85, 1.55, 2.15, title, body, color)
    voc = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(5.85), Inches(4.68), Inches(6.88), Inches(0.8))
    set_fill(voc, COLORS["blue2"])
    add_textbox(slide, 6.05, 4.83, 6.48, 0.24, "Berufliche Schulen / duale Ausbildung: praktische Ausbildung + schulische Bildung", size=14, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 0.95, 5.75, 11.7, 0.4, "Kernaussage der Arbeit: Differenzierung kann Förderung ermöglichen, legt Bildungswege aber früh fest.", size=17, bold=True, color=COLORS["navy"], align=PP_ALIGN.CENTER)
    add_footer(slide, idx, data["time"])


def add_mechanisms_slide(prs, idx, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, data["title"], data.get("section"), "pink")
    cx, cy = 6.65, 3.45
    center = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(cx - 1.1), Inches(cy - 0.65), Inches(2.2), Inches(1.3))
    set_fill(center, COLORS["white"])
    add_textbox(slide, cx - 0.92, cy - 0.2, 1.85, 0.35, "Leistungsdruck", size=18, bold=True, align=PP_ALIGN.CENTER)
    coords = [(1.05, 1.85), (8.75, 1.85), (1.05, 4.55), (8.75, 4.55)]
    fills = ["blue2", "pink2", "yellow2", "mint2"]
    for (x, y), (title, body), fill in zip(coords, data["mechanisms"], fills):
        add_card(slide, x, y, 3.85, 1.55, title, body, fill)
    add_footer(slide, idx, data["time"])


def add_interaction_slide(prs, idx, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, data["title"], data.get("section"), "lavender")
    add_textbox(slide, 1.0, 1.62, 11.3, 0.52, "Handabstimmung: Was beeinflusst euer schulisches Wohlbefinden am stärksten?", size=24, bold=True, align=PP_ALIGN.CENTER)
    options = [
        ("A", "Leistungsdruck", "Noten, Prüfungen, Vergleich", "pink"),
        ("B", "Beziehung zu Lehrkräften", "Unterstützung, Respekt, Vertrauen", "blue"),
        ("C", "Mitbestimmung", "Meinung zählt, Selbstwirksamkeit", "yellow"),
        ("D", "Lebensweltbezug", "Alltag, Beruf, Zukunft", "mint"),
    ]
    for i, (letter, title, body, color) in enumerate(options):
        x = 1.05 + (i % 2) * 5.8
        y = 2.55 + (i // 2) * 1.72
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.05), Inches(1.16))
        set_fill(card, COLORS[f"{color}2"] if color != "yellow" else COLORS["yellow2"])
        badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.2), Inches(y + 0.24), Inches(0.66), Inches(0.66))
        set_fill(badge, COLORS[color])
        add_textbox(slide, x + 0.25, y + 0.43, 0.56, 0.2, letter, size=15, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 1.05, y + 0.18, 3.75, 0.28, title, size=17, bold=True)
        add_textbox(slide, x + 1.05, y + 0.58, 3.8, 0.23, body, size=12, color=COLORS["muted"])
    note = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(5.95), Inches(9.35), Inches(0.48))
    set_fill(note, COLORS["white"])
    add_textbox(slide, 2.22, 6.08, 8.9, 0.2, "Danach: 2 kurze Stimmen einsammeln und mit den Ergebnissen der Facharbeit verknüpfen.", size=13, color=COLORS["muted"], align=PP_ALIGN.CENTER)
    add_footer(slide, idx, data["time"])


def add_wellbeing_slide(prs, idx, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, data["title"], data.get("section"), "mint")
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(1.9), Inches(5.45), Inches(3.75))
    set_fill(left, COLORS["pink2"])
    right = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.95), Inches(1.9), Inches(5.45), Inches(3.75))
    set_fill(right, COLORS["mint2"])
    add_textbox(slide, 1.25, 2.25, 4.9, 0.35, "Belastende Faktoren", size=21, bold=True, align=PP_ALIGN.CENTER)
    add_multiline(slide, 1.35, 2.9, 4.65, 1.9, [
        "Prüfungsstress und Erwartungshaltungen",
        "Vergleichskultur und Konkurrenzdenken",
        "Angst vor schulischem Versagen",
        "Schlafprobleme, Kopfschmerzen oder Konzentrationsschwierigkeiten",
    ], size=15, bullet=True, color=COLORS["navy"], spacing=4)
    add_textbox(slide, 7.25, 2.25, 4.9, 0.35, "Stärkende Faktoren", size=21, bold=True, align=PP_ALIGN.CENTER)
    add_multiline(slide, 7.35, 2.9, 4.65, 1.9, [
        "positives Schulklima",
        "unterstützende Beziehungen zu Lehrkräften",
        "Mitbestimmung und Partizipation",
        "Feedbackkultur statt nur Ergebnisfokus",
    ], size=15, bullet=True, color=COLORS["navy"], spacing=4)
    add_footer(slide, idx, data["time"])


def add_comparison_slide(prs, idx, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, data["title"], data.get("section"), "blue")
    headers = [("Deutschland", "Mehrgliedrigkeit, Noten und Prüfungen spielen zentrale Rolle", "pink2"),
               ("Skandinavische Modelle", "längeres gemeinsames Lernen; Finnland bis zur 9. Klasse", "blue2"),
               ("Angelsächsische Systeme", "umfassende Gesamtschulen, Wahlmöglichkeiten und Praxisorientierung", "yellow2")]
    for i, (title, body, fill) in enumerate(headers):
        add_card(slide, 0.92 + i * 4.15, 1.9, 3.55, 2.05, title, body, fill)
    add_textbox(slide, 1.25, 4.75, 10.8, 0.36, "Einordnung aus der Facharbeit", size=20, bold=True, align=PP_ALIGN.CENTER)
    add_multiline(slide, 1.25, 5.25, 10.8, 0.95, [
        "Länder mit späterer Leistungsdifferenzierung weisen häufig geringere Leistungsunterschiede zwischen sozialen Gruppen auf.",
        "Unterstützende Lernumgebungen und weniger Konkurrenzkultur können mit höherer Schülerzufriedenheit verbunden sein.",
        "Es gibt kein einheitliches perfektes Bildungssystem.",
    ], size=14, bullet=True, color=COLORS["navy"], spacing=2)
    add_footer(slide, idx, data["time"])


def add_proscons_slide(prs, idx, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, data["title"], data.get("section"), "yellow")
    plus = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(1.9), Inches(5.48), Inches(3.9))
    set_fill(plus, COLORS["mint2"])
    minus = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.9), Inches(5.48), Inches(3.9))
    set_fill(minus, COLORS["pink2"])
    add_textbox(slide, 1.25, 2.22, 4.85, 0.34, "Stärken", size=23, bold=True, align=PP_ALIGN.CENTER)
    add_multiline(slide, 1.45, 2.9, 4.55, 1.9, [
        "Differenzierung verschiedener Bildungswege",
        "duales Ausbildungssystem",
        "relativ niedrige Jugendarbeitslosigkeit laut OECD 2021",
        "hohe akademische Standards, besonders im Gymnasium",
    ], size=15, bullet=True, color=COLORS["navy"], spacing=3)
    add_textbox(slide, 7.2, 2.22, 4.85, 0.34, "Schwächen", size=23, bold=True, align=PP_ALIGN.CENTER)
    add_multiline(slide, 7.4, 2.9, 4.55, 1.9, [
        "frühe Aufteilung nach der Grundschule",
        "starke Kopplung von sozialer Herkunft und Bildungserfolg",
        "Leistungsdruck durch Noten, Prüfungen und Versetzung",
        "soziale Segregation zwischen Schulformen möglich",
    ], size=15, bullet=True, color=COLORS["navy"], spacing=3)
    add_footer(slide, idx, data["time"])


def add_reforms_slide(prs, idx, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, data["title"], data.get("section"), "pink")
    for i, (title, body) in enumerate(data["reforms"]):
        x = 1.0 + i * 4.05
        add_card(slide, x, 2.0, 3.45, 2.85, title, body, ["yellow2", "blue2", "lavender2"][i])
        num = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 1.25), Inches(1.55), Inches(0.86), Inches(0.86))
        set_fill(num, [COLORS["yellow"], COLORS["blue"], COLORS["lavender"]][i])
        add_textbox(slide, x + 1.31, 1.81, 0.75, 0.22, str(i + 1), size=17, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 1.1, 5.6, 11.1, 0.45, "Ziel: Bildungsqualität und Schülerwohlbefinden stärker miteinander verbinden.", size=20, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, idx, data["time"])


def add_takeaways_slide(prs, idx, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, data["title"], data.get("section"), "blue")
    add_textbox(slide, 1.1, 1.65, 11.1, 0.42, "Antwort in einem Satz:", size=18, bold=True, color=COLORS["muted"], align=PP_ALIGN.CENTER)
    answer_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(2.1), Inches(10.95), Inches(0.95))
    set_fill(answer_box, COLORS["blue2"])
    add_textbox(slide, 1.45, 2.35, 10.45, 0.32, "Das Wohlbefinden hängt nicht nur von einzelnen Schülerinnen und Schülern ab, sondern stark von schulischen Strukturen.", size=18, bold=True, align=PP_ALIGN.CENTER)
    for i, text in enumerate(data["takeaways"]):
        y = 3.55 + i * 0.66
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(1.55), Inches(y + 0.05), Inches(0.23), Inches(0.23))
        set_fill(dot, [COLORS["pink"], COLORS["yellow"], COLORS["mint"], COLORS["lavender"]][i])
        add_textbox(slide, 1.95, y, 9.8, 0.35, text, size=15, color=COLORS["navy"])
    add_footer(slide, idx, data["time"])


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs, SLIDES[0])
    add_cards_slide(prs, 2, SLIDES[1], "yellow")
    add_leitfrage_slide(prs, 3, SLIDES[2])
    add_cards_slide(prs, 4, SLIDES[3], "blue")
    add_functions_slide(prs, 5, SLIDES[4])
    add_flow_slide(prs, 6, SLIDES[5])
    add_mechanisms_slide(prs, 7, SLIDES[6])
    add_interaction_slide(prs, 8, SLIDES[7])
    add_wellbeing_slide(prs, 9, SLIDES[8])
    add_cards_slide(prs, 10, SLIDES[9], "mint")
    add_comparison_slide(prs, 11, SLIDES[10])
    add_proscons_slide(prs, 12, SLIDES[11])
    add_reforms_slide(prs, 13, SLIDES[12])
    add_takeaways_slide(prs, 14, SLIDES[13])

    prs.save(PPTX_PATH)
    apply_motion_to_presentation(PPTX_PATH)
    copy2(PPTX_PATH, KEYNOTE_PPTX_PATH)
    apply_keynote_safe_transitions(KEYNOTE_PPTX_PATH)


def qn(namespace, tag):
    return f"{{{namespace}}}{tag}"


def p_el(tag, attrs=None, children=None):
    el = etree.Element(qn(P_NS, tag), attrs or {})
    for child in children or []:
        el.append(child)
    return el


def shape_text(shape):
    paragraphs = []
    for paragraph in shape.xpath(".//a:p", namespaces=NS):
        text = "".join(t.text or "" for t in paragraph.xpath(".//a:t", namespaces=NS)).strip()
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs).strip()


def slide_shapes(root):
    shapes = {}
    for shape in root.xpath(".//p:sp", namespaces=NS):
        c_nv_pr = shape.find(".//p:cNvPr", namespaces=NS)
        if c_nv_pr is None:
            continue
        text = shape_text(shape)
        if text:
            shapes.setdefault(text, []).append(c_nv_pr.get("id"))
    return shapes


def find_shape(shapes, text):
    matches = shapes.get(text)
    if matches:
        return matches[0]
    for shape_text_value, ids in shapes.items():
        if shape_text_value.replace("\n", " ") == text:
            return ids[0]
    for shape_text_value, ids in shapes.items():
        if shape_text_value.startswith(text):
            return ids[0]
    return None


def paragraph_indices(root, spid):
    shape = root.xpath(f".//p:sp[p:nvSpPr/p:cNvPr[@id='{spid}']]", namespaces=NS)
    if not shape:
        return []
    indices = []
    for idx, paragraph in enumerate(shape[0].xpath(".//a:p", namespaces=NS)):
        text = "".join(t.text or "" for t in paragraph.xpath(".//a:t", namespaces=NS)).strip()
        if text:
            indices.append(idx)
    return indices


def add_transition(root, effect="fade", direction=None):
    for old in root.xpath("./p:transition | ./mc:AlternateContent[.//p:transition]", namespaces=NS):
        old.getparent().remove(old)

    root.set(qn(MC_NS, "Ignorable"), "p14")

    alternate = etree.Element(qn(MC_NS, "AlternateContent"), nsmap={"mc": MC_NS, "p14": P14_NS})
    choice = etree.SubElement(alternate, qn(MC_NS, "Choice"), Requires="p14")
    choice_transition = etree.SubElement(choice, qn(P_NS, "transition"), spd="med")
    choice_transition.set(qn(P14_NS, "dur"), "700")
    if effect == "wipe":
        attrs = {"dir": direction or "r"}
        etree.SubElement(choice_transition, qn(P_NS, "wipe"), attrs)
    else:
        etree.SubElement(choice_transition, qn(P_NS, "fade"))
    fallback = etree.SubElement(alternate, qn(MC_NS, "Fallback"))
    fallback_transition = etree.SubElement(fallback, qn(P_NS, "transition"), spd="med")
    if effect == "wipe":
        attrs = {"dir": direction or "r"}
        etree.SubElement(fallback_transition, qn(P_NS, "wipe"), attrs)
    else:
        etree.SubElement(fallback_transition, qn(P_NS, "fade"))

    timing = root.find(qn(P_NS, "timing"))
    if timing is not None:
        root.insert(root.index(timing), alternate)
        return
    clr_map = root.find(qn(P_NS, "clrMapOvr"))
    insert_at = root.index(clr_map) + 1 if clr_map is not None else 1
    root.insert(insert_at, alternate)


def add_fade_transition(root):
    add_transition(root, "fade")


def text_target(spid, paragraph_idx=None):
    sp_tgt = p_el("spTgt", {"spid": str(spid)})
    if paragraph_idx is not None:
        tx_el = p_el("txEl")
        tx_el.append(p_el("pRg", {"st": str(paragraph_idx), "end": str(paragraph_idx)}))
        sp_tgt.append(tx_el)
    tgt_el = p_el("tgtEl")
    tgt_el.append(sp_tgt)
    return tgt_el


def anim_scale(effect_id, spid, duration=420):
    anim = p_el("animScale")
    c_bhvr = p_el("cBhvr")
    c_bhvr.append(p_el("cTn", {"id": str(effect_id), "dur": str(duration), "fill": "hold"}))
    c_bhvr.append(text_target(spid))
    attr_names = p_el("attrNameLst")
    attr_x = p_el("attrName")
    attr_x.text = "ScaleX"
    attr_y = p_el("attrName")
    attr_y.text = "ScaleY"
    attr_names.extend([attr_x, attr_y])
    c_bhvr.append(attr_names)
    anim.append(c_bhvr)
    anim.append(p_el("from", {"x": "96000", "y": "96000"}))
    anim.append(p_el("to", {"x": "100000", "y": "100000"}))
    return anim


def set_visible(effect_id, spid, paragraph_idx=None):
    setter = p_el("set")
    c_bhvr = p_el("cBhvr")
    c_bhvr.append(p_el("cTn", {"id": str(effect_id), "dur": "1", "fill": "hold"}))
    c_bhvr.append(text_target(spid, paragraph_idx))
    attr_names = p_el("attrNameLst")
    attr_name = p_el("attrName")
    attr_name.text = "style.visibility"
    attr_names.append(attr_name)
    c_bhvr.append(attr_names)
    setter.append(c_bhvr)
    to = p_el("to")
    to.append(p_el("strVal", {"val": "visible"}))
    setter.append(to)
    return setter


def entrance_preset(effect):
    if effect == "wipe":
        return {"presetID": "2819", "presetClass": "entr", "presetSubtype": "0"}
    if effect == "zoom":
        return {"presetID": "22", "presetClass": "entr", "presetSubtype": "0"}
    return {"presetID": "10", "presetClass": "entr", "presetSubtype": "0"}


def anim_effect(effect_id, spid, delay, duration=450, paragraph_idx=None, node_type="afterEffect", effect="fade"):
    par = p_el("par")
    ctn_attrs = {"id": str(effect_id), "fill": "hold", "nodeType": node_type}
    ctn_attrs.update(entrance_preset(effect))
    ctn = p_el("cTn", ctn_attrs)
    st_cond_lst = p_el("stCondLst")
    st_cond_lst.append(p_el("cond", {"delay": str(delay)}))
    ctn.append(st_cond_lst)
    child_tn_lst = p_el("childTnLst")

    # Office documents Wipe with directional values like wipe(right); Keynote is
    # more likely to import these standard filter names than custom aliases.
    filter_name = "wipe(right)" if effect == "wipe" else "fade"
    child_tn_lst.append(set_visible(effect_id + 1, spid, paragraph_idx))
    anim = p_el("animEffect", {"transition": "in", "filter": filter_name})
    c_bhvr = p_el("cBhvr")
    c_bhvr.append(p_el("cTn", {"id": str(effect_id + 2), "dur": str(duration), "fill": "hold"}))
    c_bhvr.append(text_target(spid, paragraph_idx))
    anim.append(c_bhvr)
    child_tn_lst.append(anim)

    if effect == "zoom" and paragraph_idx is None:
        child_tn_lst.append(anim_scale(effect_id + 3, spid, duration=duration))

    ctn.append(child_tn_lst)
    par.append(ctn)
    return par


def slide_trigger_conditions(tag):
    cond_lst = p_el(tag)
    cond = p_el("cond", {"evt": "onNext" if tag == "nextCondLst" else "onPrev", "delay": "0"})
    tgt = p_el("tgtEl")
    tgt.append(p_el("sldTgt"))
    cond.append(tgt)
    cond_lst.append(cond)
    return cond_lst


def build_timing(root, animation_groups):
    timing = root.find(qn(P_NS, "timing"))
    if timing is not None:
        root.remove(timing)

    timing = p_el("timing")
    tn_lst = p_el("tnLst")
    seq = p_el("seq", {"concurrent": "1", "nextAc": "seek"})
    main_ctn = p_el("cTn", {"id": "1", "dur": "indefinite", "nodeType": "mainSeq"})
    st_cond_lst = p_el("stCondLst")
    st_cond_lst.append(p_el("cond", {"delay": "0"}))
    main_ctn.append(st_cond_lst)
    child_tn_lst = p_el("childTnLst")

    effect_id = 2
    paragraph_build_shapes = set()
    for group in animation_groups:
        delay = group["delay"]
        duration = group.get("duration", 450)
        targets = group["targets"]
        for target_idx, target in enumerate(targets):
            spid = target["spid"]
            paragraph_idx = target.get("paragraph")
            effect = target.get("effect", group.get("effect", "fade"))
            if paragraph_idx is not None:
                paragraph_build_shapes.add(spid)
            node_type = "afterEffect" if target_idx == 0 else "withEffect"
            child_tn_lst.append(
                anim_effect(
                    effect_id,
                    spid,
                    delay=delay,
                    duration=duration,
                    paragraph_idx=paragraph_idx,
                    node_type=node_type,
                    effect=effect,
                )
            )
            effect_id += 4

    main_ctn.append(child_tn_lst)
    seq.append(main_ctn)
    seq.append(slide_trigger_conditions("prevCondLst"))
    seq.append(slide_trigger_conditions("nextCondLst"))
    tn_lst.append(seq)
    timing.append(tn_lst)

    if paragraph_build_shapes:
        bld_lst = p_el("bldLst")
        for spid in sorted(paragraph_build_shapes, key=int):
            bld_lst.append(p_el("bldP", {"spid": spid, "grpId": "0", "build": "p"}))
        timing.append(bld_lst)

    root.append(timing)


def group_for_texts(shapes, texts):
    targets = []
    for text in texts:
        spid = find_shape(shapes, text)
        if spid:
            targets.append({"spid": spid})
    return targets


def groups_from_card_data(shapes, cards, start_delay):
    groups = []
    delay = start_delay
    for title, body, _ in cards:
        targets = group_for_texts(shapes, [title, body])
        if targets:
            groups.append({"delay": delay, "duration": 430, "effect": "zoom", "targets": targets})
            delay += 270
    return groups


def groups_from_paragraphs(root, spid, start_delay, step=230, duration=350):
    groups = []
    delay = start_delay
    for idx in paragraph_indices(root, spid):
        groups.append({"delay": delay, "duration": duration, "effect": "wipe", "targets": [{"spid": spid, "paragraph": idx}]})
        delay += step
    return groups


def animation_plan(slide_no, root, data):
    shapes = slide_shapes(root)
    groups = []

    title_id = find_shape(shapes, data["title"])
    if title_id:
        groups.append({"delay": 0, "duration": 500, "targets": [{"spid": title_id}]})

    start = 560

    if slide_no == 1:
        for text in [data["subtitle"], "Facharbeit von Svea Timphus", "Leitfrage, zentrale Ergebnisse und Ausblick"]:
            targets = group_for_texts(shapes, [text])
            if targets:
                effect = "wipe" if text == data["subtitle"] else "fade"
                groups.append({"delay": start, "duration": 450, "effect": effect, "targets": targets})
                start += 280
        return groups

    if data.get("interaction"):
        question = "Handabstimmung: Was beeinflusst euer schulisches Wohlbefinden am stärksten?"
        targets = group_for_texts(shapes, [question])
        if targets:
            groups.append({"delay": start, "duration": 450, "effect": "wipe", "targets": targets})
            start += 390
        options = [
            ["A", "Leistungsdruck", "Noten, Prüfungen, Vergleich"],
            ["B", "Beziehung zu Lehrkräften", "Unterstützung, Respekt, Vertrauen"],
            ["C", "Mitbestimmung", "Meinung zählt, Selbstwirksamkeit"],
            ["D", "Lebensweltbezug", "Alltag, Beruf, Zukunft"],
        ]
        for option in options:
            targets = group_for_texts(shapes, option)
            if targets:
                groups.append({"delay": start, "duration": 360, "effect": "zoom", "targets": targets})
                start += 300
        note_targets = group_for_texts(shapes, ["Danach: 2 kurze Stimmen einsammeln und mit den Ergebnissen der Facharbeit verknüpfen."])
        if note_targets:
            groups.append({"delay": start + 120, "duration": 350, "effect": "fade", "targets": note_targets})
        return groups

    if "takeaways" in data:
        answer_targets = group_for_texts(
            shapes,
            [
                "Antwort in einem Satz:",
                "Das Wohlbefinden hängt nicht nur von einzelnen Schülerinnen und Schülern ab, sondern stark von schulischen Strukturen.",
            ],
        )
        if answer_targets:
            groups.append({"delay": start, "duration": 500, "effect": "fade", "targets": answer_targets})
            start += 430
        for takeaway in data["takeaways"]:
            targets = group_for_texts(shapes, [takeaway])
            if targets:
                groups.append({"delay": start, "duration": 360, "effect": "wipe", "targets": targets})
                start += 260
        return groups

    if "cards" in data:
        groups.extend(groups_from_card_data(shapes, data["cards"], start))
        return groups

    if "bullets" in data:
        quote_targets = group_for_texts(shapes, [f"\"{data['quote']}\""])
        if quote_targets:
            groups.append({"delay": start, "duration": 450, "effect": "fade", "targets": quote_targets})
            start += 360
        bullet_shape = find_shape(shapes, "\n".join(f"- {bullet}" for bullet in data["bullets"]))
        if bullet_shape:
            groups.extend(groups_from_paragraphs(root, bullet_shape, start, step=230, duration=330))
        return groups

    if "functions" in data:
        for title, body in data["functions"]:
            targets = group_for_texts(shapes, [title, body])
            if targets:
                groups.append({"delay": start, "duration": 420, "effect": "zoom", "targets": targets})
                start += 280
        return groups

    if data.get("flow"):
        sequence = [
            ["Grundschule", "gemeinsamer Start"],
            ["oft nach Klasse 4"],
            ["Gymnasium", "Abitur, stärker theoretisch, hoher Prüfungsdruck möglich"],
            ["Realschule", "Mittlerer Abschluss, Verbindung von Allgemeinbildung und Praxis"],
            ["Hauptschule", "stärker praktisch, teils mit Stigmatisierung verbunden"],
            ["Gesamtschule", "längeres gemeinsames Lernen, mehrere Abschlüsse"],
            ["Berufliche Schulen / duale Ausbildung: praktische Ausbildung + schulische Bildung"],
            ["Kernaussage der Arbeit: Differenzierung kann Förderung ermöglichen, legt Bildungswege aber früh fest."],
        ]
        for item in sequence:
            targets = group_for_texts(shapes, item)
            if targets:
                effect = "wipe" if len(item) == 1 else "zoom"
                groups.append({"delay": start, "duration": 360, "effect": effect, "targets": targets})
                start += 230
        return groups

    if "mechanisms" in data:
        center_targets = group_for_texts(shapes, ["Leistungsdruck"])
        if center_targets:
            groups.append({"delay": start, "duration": 420, "effect": "zoom", "targets": center_targets})
            start += 300
        for title, body in data["mechanisms"]:
            targets = group_for_texts(shapes, [title, body])
            if targets:
                groups.append({"delay": start, "duration": 380, "effect": "zoom", "targets": targets})
                start += 260
        return groups

    if data.get("wellbeing"):
        for heading, bullets in [
            ("Belastende Faktoren", [
                "Prüfungsstress und Erwartungshaltungen",
                "Vergleichskultur und Konkurrenzdenken",
                "Angst vor schulischem Versagen",
                "Schlafprobleme, Kopfschmerzen oder Konzentrationsschwierigkeiten",
            ]),
            ("Stärkende Faktoren", [
                "positives Schulklima",
                "unterstützende Beziehungen zu Lehrkräften",
                "Mitbestimmung und Partizipation",
                "Feedbackkultur statt nur Ergebnisfokus",
            ]),
        ]:
            heading_targets = group_for_texts(shapes, [heading])
            if heading_targets:
                groups.append({"delay": start, "duration": 350, "effect": "fade", "targets": heading_targets})
                start += 210
            bullet_shape = find_shape(shapes, "\n".join(f"- {bullet}" for bullet in bullets))
            if bullet_shape:
                groups.extend(groups_from_paragraphs(root, bullet_shape, start, step=170, duration=300))
                start += len(paragraph_indices(root, bullet_shape)) * 170
        return groups

    if data.get("comparison"):
        comparison_cards = [
            ("Deutschland", "Mehrgliedrigkeit, Noten und Prüfungen spielen zentrale Rolle", ""),
            ("Skandinavische Modelle", "längeres gemeinsames Lernen; Finnland bis zur 9. Klasse", ""),
            ("Angelsächsische Systeme", "umfassende Gesamtschulen, Wahlmöglichkeiten und Praxisorientierung", ""),
        ]
        groups.extend(groups_from_card_data(shapes, comparison_cards, start))
        start += 820
        heading_targets = group_for_texts(shapes, ["Einordnung aus der Facharbeit"])
        if heading_targets:
            groups.append({"delay": start, "duration": 340, "effect": "fade", "targets": heading_targets})
            start += 230
        bullets = [
            "Länder mit späterer Leistungsdifferenzierung weisen häufig geringere Leistungsunterschiede zwischen sozialen Gruppen auf.",
            "Unterstützende Lernumgebungen und weniger Konkurrenzkultur können mit höherer Schülerzufriedenheit verbunden sein.",
            "Es gibt kein einheitliches perfektes Bildungssystem.",
        ]
        bullet_shape = find_shape(shapes, "\n".join(f"- {bullet}" for bullet in bullets))
        if bullet_shape:
            groups.extend(groups_from_paragraphs(root, bullet_shape, start, step=210, duration=300))
        return groups

    if data.get("proscons"):
        for heading, bullets in [
            ("Stärken", [
                "Differenzierung verschiedener Bildungswege",
                "duales Ausbildungssystem",
                "relativ niedrige Jugendarbeitslosigkeit laut OECD 2021",
                "hohe akademische Standards, besonders im Gymnasium",
            ]),
            ("Schwächen", [
                "frühe Aufteilung nach der Grundschule",
                "starke Kopplung von sozialer Herkunft und Bildungserfolg",
                "Leistungsdruck durch Noten, Prüfungen und Versetzung",
                "soziale Segregation zwischen Schulformen möglich",
            ]),
        ]:
            heading_targets = group_for_texts(shapes, [heading])
            if heading_targets:
                groups.append({"delay": start, "duration": 340, "effect": "fade", "targets": heading_targets})
                start += 190
            bullet_shape = find_shape(shapes, "\n".join(f"- {bullet}" for bullet in bullets))
            if bullet_shape:
                groups.extend(groups_from_paragraphs(root, bullet_shape, start, step=155, duration=290))
                start += len(paragraph_indices(root, bullet_shape)) * 155
        return groups

    if "reforms" in data:
        for title, body in data["reforms"]:
            targets = group_for_texts(shapes, [title, body])
            if targets:
                groups.append({"delay": start, "duration": 420, "effect": "zoom", "targets": targets})
                start += 300
        goal_targets = group_for_texts(shapes, ["Ziel: Bildungsqualität und Schülerwohlbefinden stärker miteinander verbinden."])
        if goal_targets:
            groups.append({"delay": start + 100, "duration": 380, "effect": "fade", "targets": goal_targets})
        return groups

    return groups


def apply_motion_to_presentation(pptx_path):
    tmp_path = pptx_path.with_suffix(".animated.tmp.pptx")
    with ZipFile(pptx_path, "r") as source, ZipFile(tmp_path, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("ppt/slides/slide") and info.filename.endswith(".xml"):
                slide_no = int(info.filename.rsplit("slide", 1)[1].split(".xml", 1)[0])
                if 1 <= slide_no <= len(SLIDES):
                    root = etree.fromstring(data)
                    add_fade_transition(root)
                    build_timing(root, animation_plan(slide_no, root, SLIDES[slide_no - 1]))
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(info, data)
    tmp_path.replace(pptx_path)


def apply_keynote_safe_transitions(pptx_path):
    # Keynote on iPad often simplifies imported PowerPoint object builds. This
    # variant therefore adds visible, conservative slide-level variety using
    # effects Keynote commonly maps reliably: Fade and Wipe.
    transition_cycle = [
        ("fade", None),
        ("wipe", "r"),
        ("fade", None),
        ("wipe", "u"),
        ("fade", None),
        ("wipe", "r"),
        ("fade", None),
        ("wipe", "d"),
        ("fade", None),
        ("wipe", "r"),
        ("fade", None),
        ("wipe", "u"),
        ("fade", None),
        ("wipe", "r"),
    ]
    tmp_path = pptx_path.with_suffix(".keynote.tmp.pptx")
    with ZipFile(pptx_path, "r") as source, ZipFile(tmp_path, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("ppt/slides/slide") and info.filename.endswith(".xml"):
                slide_no = int(info.filename.rsplit("slide", 1)[1].split(".xml", 1)[0])
                if 1 <= slide_no <= len(transition_cycle):
                    root = etree.fromstring(data)
                    effect, direction = transition_cycle[slide_no - 1]
                    add_transition(root, effect, direction)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(info, data)
    tmp_path.replace(pptx_path)


def build_guide():
    lines = [
        "# Sprechleitfaden zur Präsentation",
        "",
        "**Thema:** Wie Schule wirkt: strukturelle Unterschiede und ihr Effekt für das Wohlbefinden der Schüler",
        "",
        "**Ziel:** ca. 15 Minuten Vortrag im Seminarfachkurs Glück. Alle inhaltlichen Punkte stammen aus der Facharbeit; die Abstimmung ist nur eine Aktivierung des Publikums.",
        "",
        "## Ablauf und Sprechimpulse",
        "",
    ]

    for idx, slide in enumerate(SLIDES, start=1):
        lines.append(f"### Folie {idx}: {slide['title']} ({slide['time']})")
        lines.append("")
        for item in slide.get("speaker", []):
            lines.append(f"- {item}")
        if slide.get("interaction"):
            lines.extend([
                "",
                "**Publikumseinbindung:**",
                "- Frage stellen: \"Was beeinflusst euer schulisches Wohlbefinden am stärksten?\"",
                "- Handzeichen für A bis D sammeln.",
                "- Optional zwei kurze Nachfragen: \"Warum gerade dieser Punkt?\"",
                "- Danach überleiten: \"Genau diese Faktoren untersucht die Facharbeit beim Zusammenhang von Schule und Wohlbefinden.\"",
            ])
        lines.append("")

    lines.extend([
        "## Kurzer Schluss-Satz",
        "",
        "Die Facharbeit zeigt, dass Schule nicht nur Wissen vermittelt, sondern durch Struktur, Bewertung, Beziehungen und Lebensweltbezug stark mitbestimmt, wie Schülerinnen und Schüler Schule erleben und wie wohl sie sich dort fühlen.",
        "",
        "## Quellenbasis laut Facharbeit",
        "",
        "- Autorengruppe Bildungsberichterstattung 2022",
        "- OECD 2019 und OECD 2021",
        "- WHO 2020 / HBSC",
        "- Robert Koch-Institut / KiGGS",
        "- KMK 2023 sowie weitere im Literaturverzeichnis der Facharbeit genannte Quellen",
    ])

    GUIDE_PATH.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    build_presentation()
    build_guide()
    print(f"Created {PPTX_PATH}")
    print(f"Created {KEYNOTE_PPTX_PATH}")
    print(f"Created {GUIDE_PATH}")
