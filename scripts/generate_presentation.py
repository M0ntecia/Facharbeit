from __future__ import annotations

import os
import shutil
import zipfile
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path("presentation")
PPTX_PATH = OUT_DIR / "Wie_Schule_wirkt_Keynote_Style.pptx"
NOTES_PATH = OUT_DIR / "Sprechnotizen.md"

SLIDE_W = 13.333
SLIDE_H = 7.5


class C:
    BLUE = RGBColor(0xCF, 0xEF, 0xFF)
    PINK = RGBColor(0xFF, 0xD9, 0xE8)
    YELLOW = RGBColor(0xFF, 0xF7, 0xC7)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    DARK = RGBColor(0x36, 0x38, 0x3D)
    MID = RGBColor(0x6E, 0x72, 0x78)
    LIGHT = RGBColor(0xF7, 0xFB, 0xFF)
    GREEN = RGBColor(0xD8, 0xF7, 0xDF)
    LAV = RGBColor(0xEA, 0xE6, 0xFF)
    SHADOW = RGBColor(0xE7, 0xEC, 0xF1)


FONT_HEAD = "Aptos Display"
FONT_BODY = "Aptos"


@dataclass
class SlideInfo:
    title: str
    transition: str
    animation_ids: list[int]
    notes: str


slide_infos: list[SlideInfo] = []


def i(value: float):
    return Inches(value)


def add_bg(slide, accent=C.BLUE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C.WHITE
    # Soft corner blobs for an Apple-Keynote-like pastel canvas.
    blob1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, i(-1.1), i(-1.0), i(4.2), i(3.0))
    blob1.fill.solid()
    blob1.fill.fore_color.rgb = accent
    blob1.line.fill.background()
    blob2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, i(10.7), i(5.6), i(3.3), i(2.4))
    blob2.fill.solid()
    blob2.fill.fore_color.rgb = C.YELLOW
    blob2.line.fill.background()


def add_shadow(slide, x, y, w, h, radius=MSO_SHAPE.ROUNDED_RECTANGLE):
    shadow = slide.shapes.add_shape(radius, i(x + 0.06), i(y + 0.07), i(w), i(h))
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = C.SHADOW
    shadow.line.fill.background()
    return shadow


def add_card(slide, x, y, w, h, fill=C.WHITE, line=None, radius=MSO_SHAPE.ROUNDED_RECTANGLE):
    add_shadow(slide, x, y, w, h, radius)
    card = slide.shapes.add_shape(radius, i(x), i(y), i(w), i(h))
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    if line:
        card.line.color.rgb = line
        card.line.width = Pt(1.0)
    else:
        card.line.fill.background()
    return card


def set_text_style(paragraph, size=24, bold=False, color=C.DARK, align=None):
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        run.font.name = FONT_BODY
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def text_box(
    slide,
    x,
    y,
    w,
    h,
    text,
    size=24,
    bold=False,
    color=C.DARK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font=FONT_BODY,
    animate=False,
    margin=0.05,
):
    shape = slide.shapes.add_textbox(i(x), i(y), i(w), i(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = i(margin)
    tf.margin_right = i(margin)
    tf.margin_top = i(margin)
    tf.margin_bottom = i(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if animate:
        current_animation_ids().append(shape.shape_id)
    return shape


def bullet_list(slide, x, y, w, h, bullets, size=22, color=C.DARK, animate=True, line_spacing=1.1):
    box = slide.shapes.add_textbox(i(x), i(y), i(w), i(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = i(0.08)
    tf.margin_right = i(0.08)
    tf.margin_top = i(0.02)
    tf.margin_bottom = i(0.02)
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = FONT_BODY
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(8 * line_spacing)
        p._p.get_or_add_pPr().set("marL", "228600")
        p._p.get_or_add_pPr().set("indent", "-127000")
        p._p.get_or_add_pPr().append(etree.Element("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", char="•"))
    if animate:
        current_animation_ids().append(box.shape_id)
    return box


def chip(slide, x, y, w, text, fill, size=17):
    shp = add_card(slide, x, y, w, 0.42, fill=fill, radius=MSO_SHAPE.ROUNDED_RECTANGLE)
    text_box(slide, x + 0.08, y + 0.08, w - 0.16, 0.25, text, size=size, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return shp


def title(slide, text, subtitle=None, x=0.7, y=0.55, w=7.8, title_size=42):
    text_box(slide, x, y, w, 0.7, text, size=title_size, bold=True, font=FONT_HEAD, animate=True)
    if subtitle:
        text_box(slide, x, y + 0.82, w, 0.35, subtitle, size=18, color=C.MID, animate=True)


def current_animation_ids() -> list[int]:
    return slide_infos[-1].animation_ids


def begin_slide(prs, title_text, transition, accent=C.BLUE, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_infos.append(SlideInfo(title_text, transition, [], notes))
    add_bg(slide, accent)
    return slide


def add_notes(slide, notes):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.clear()
    text_frame.text = notes


def line(slide, x1, y1, x2, y2, color=C.MID, width=1.7):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, i(x1), i(y1), i(x2), i(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    return ln


def draw_person(slide, x, y, scale=1.0, shirt=C.BLUE, mood="smile"):
    head = slide.shapes.add_shape(MSO_SHAPE.OVAL, i(x), i(y), i(0.28 * scale), i(0.28 * scale))
    head.fill.solid()
    head.fill.fore_color.rgb = C.YELLOW
    head.line.fill.background()
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, i(x - 0.02 * scale), i(y + 0.27 * scale), i(0.32 * scale), i(0.45 * scale))
    body.fill.solid()
    body.fill.fore_color.rgb = shirt
    body.line.fill.background()
    eye1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, i(x + 0.08 * scale), i(y + 0.10 * scale), i(0.025 * scale), i(0.025 * scale))
    eye2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, i(x + 0.18 * scale), i(y + 0.10 * scale), i(0.025 * scale), i(0.025 * scale))
    for eye in (eye1, eye2):
        eye.fill.solid()
        eye.fill.fore_color.rgb = C.DARK
        eye.line.fill.background()
    if mood == "stress":
        mouth = line(slide, x + 0.09 * scale, y + 0.20 * scale, x + 0.21 * scale, y + 0.18 * scale, C.DARK, 1.2)
        bolt = text_box(slide, x + 0.26 * scale, y - 0.08 * scale, 0.18 * scale, 0.18 * scale, "!", size=13 * scale, bold=True, color=C.PINK)
        bolt.rotation = -15
    else:
        mouth = line(slide, x + 0.09 * scale, y + 0.18 * scale, x + 0.21 * scale, y + 0.18 * scale, C.DARK, 1.2)
    return mouth


def draw_school(slide, x, y, w, h):
    base = add_card(slide, x, y + 0.7, w, h - 0.7, fill=C.BLUE)
    roof = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, i(x + 0.3), i(y), i(w - 0.6), i(1.1))
    roof.fill.solid()
    roof.fill.fore_color.rgb = C.PINK
    roof.line.fill.background()
    door = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, i(x + w / 2 - 0.32), i(y + h - 0.82), i(0.64), i(0.72))
    door.fill.solid()
    door.fill.fore_color.rgb = C.WHITE
    door.line.fill.background()
    for row_y in (y + 1.02, y + 1.58):
        for col in range(4):
            win = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, i(x + 0.42 + col * 0.68), i(row_y), i(0.38), i(0.28))
            win.fill.solid()
            win.fill.fore_color.rgb = C.WHITE
            win.line.fill.background()
    text_box(slide, x + 1.0, y + 0.75, 1.25, 0.3, "SCHULE", size=13, bold=True, align=PP_ALIGN.CENTER)
    return base


def draw_scale(slide, x, y):
    line(slide, x + 1.6, y + 0.5, x + 1.6, y + 2.4, C.DARK, 2)
    line(slide, x + 0.35, y + 0.72, x + 2.85, y + 0.72, C.DARK, 2)
    line(slide, x + 0.55, y + 0.72, x + 0.28, y + 1.35, C.MID, 1.3)
    line(slide, x + 1.0, y + 0.72, x + 1.27, y + 1.35, C.MID, 1.3)
    line(slide, x + 2.2, y + 0.72, x + 1.92, y + 1.35, C.MID, 1.3)
    line(slide, x + 2.65, y + 0.72, x + 2.92, y + 1.35, C.MID, 1.3)
    left = slide.shapes.add_shape(MSO_SHAPE.ARC, i(x + 0.18), i(y + 1.1), i(1.22), i(0.62))
    right = slide.shapes.add_shape(MSO_SHAPE.ARC, i(x + 1.8), i(y + 1.1), i(1.22), i(0.62))
    for arc in (left, right):
        arc.line.color.rgb = C.DARK
        arc.line.width = Pt(2)
    text_box(slide, x - 0.1, y + 1.75, 1.6, 0.35, "Bildung", size=16, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, x + 1.72, y + 1.75, 1.6, 0.35, "Wohlbefinden", size=16, bold=True, align=PP_ALIGN.CENTER)


def draw_stressed_student(slide, x, y):
    add_card(slide, x, y, 3.0, 3.7, fill=C.LIGHT)
    draw_person(slide, x + 1.25, y + 0.45, scale=2.1, shirt=C.PINK, mood="stress")
    text_box(slide, x + 0.55, y + 2.9, 1.9, 0.3, "Prüfung", size=17, bold=True, align=PP_ALIGN.CENTER)
    for dx, dy, txt in [(0.25, 0.55, "Note"), (2.2, 0.75, "!"), (0.45, 2.25, "Test"), (2.22, 2.05, "?")]:
        chip(slide, x + dx, y + dy, 0.58, txt, C.YELLOW, size=12)


def draw_dialogue(slide, x, y):
    add_card(slide, x, y, 3.3, 2.55, fill=C.WHITE)
    draw_person(slide, x + 0.48, y + 0.55, scale=1.55, shirt=C.BLUE)
    draw_person(slide, x + 2.2, y + 0.58, scale=1.35, shirt=C.PINK)
    bubble1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, i(x + 1.18), i(y + 0.34), i(0.92), i(0.42))
    bubble1.fill.solid()
    bubble1.fill.fore_color.rgb = C.YELLOW
    bubble1.line.fill.background()
    bubble2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, i(x + 1.62), i(y + 1.05), i(0.78), i(0.38))
    bubble2.fill.solid()
    bubble2.fill.fore_color.rgb = C.BLUE
    bubble2.line.fill.background()


def draw_future_arrow(slide, x, y):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, i(x), i(y), i(4.5), i(1.0))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = C.BLUE
    arrow.line.fill.background()
    text_box(slide, x + 1.25, y + 0.27, 2.0, 0.3, "Zukunft", size=22, bold=True, align=PP_ALIGN.CENTER)


def slide_1(prs):
    notes = (
        "Ich beginne mit der Grundidee meiner Facharbeit: Schule ist nicht nur ein Ort, an dem "
        "Wissen vermittelt und Leistungen bewertet werden. Sie ist auch ein Lebensraum, in dem "
        "Schülerinnen und Schüler einen großen Teil ihres Alltags verbringen. Deshalb wirkt Schule "
        "auf Motivation, Selbstwertgefühl und Zukunftsperspektiven. In meiner Arbeit untersuche ich, "
        "wie strukturelle Unterschiede im deutschen Schulsystem mit dem Wohlbefinden zusammenhängen. "
        "Dabei geht es nicht darum, Schule grundsätzlich schlechtzureden, sondern die Spannung "
        "zwischen Bildungsauftrag und Lebensrealität sichtbar zu machen. Von dieser Ausgangsidee "
        "führt die nächste Folie zur Frage, warum das Thema gerade im Seminarfachkurs Glück wichtig ist."
    )
    slide = begin_slide(prs, "Titel", "Fade", C.BLUE, notes)
    text_box(slide, 0.72, 0.78, 6.9, 0.72, "Wie Schule wirkt:", size=46, bold=True, font=FONT_HEAD, animate=True)
    text_box(
        slide,
        0.76,
        1.58,
        6.65,
        1.25,
        "Strukturelle Unterschiede und ihr Effekt auf das Wohlbefinden von Schülerinnen und Schülern",
        size=27,
        bold=True,
        font=FONT_HEAD,
        animate=True,
    )
    pill = add_card(slide, 0.78, 3.12, 2.65, 0.48, fill=C.PINK)
    text_box(slide, 0.94, 3.24, 2.3, 0.22, "Seminarfachkurs Glück", size=17, bold=True, align=PP_ALIGN.CENTER, animate=True)
    draw_school(slide, 8.45, 1.15, 3.25, 2.9)
    draw_person(slide, 8.0, 4.48, 1.65, C.PINK, "smile")
    draw_person(slide, 9.8, 4.35, 1.45, C.BLUE, "stress")
    draw_person(slide, 11.05, 4.65, 1.25, C.YELLOW, "smile")
    add_notes(slide, notes)


def slide_2(prs):
    notes = (
        "Die Relevanz des Themas ergibt sich daraus, dass Schule den Alltag junger Menschen stark "
        "prägt. In der Facharbeit wird Schule als Institution beschrieben, die Wissen vermittelt, "
        "aber gleichzeitig auch bewertet, vergleicht und selektiert. Damit hängen Bildungswege und "
        "Zukunftschancen eng mit schulischen Erfahrungen zusammen. Zugleich wird in der Arbeit auf "
        "die öffentliche Diskussion über Leistungsdruck, psychische Belastungen und Praxisnähe "
        "verwiesen. Für einen Kurs zum Thema Glück ist daran besonders interessant: Wohlbefinden ist "
        "nicht nur ein privates Gefühl, sondern ein Bestandteil gelingender Bildung. Daraus entsteht "
        "die Leitfrage, die den weiteren Vortrag strukturiert."
    )
    slide = begin_slide(prs, "Warum ist das Thema wichtig?", "Zoom", C.PINK, notes)
    title(slide, "Warum ist das Thema wichtig?", subtitle="Schule zwischen Bildungsauftrag und Lebensrealität")
    draw_scale(slide, 8.75, 1.0)
    bullet_list(
        slide,
        0.88,
        2.08,
        5.2,
        2.3,
        [
            "Schule prägt Alltag",
            "Leistungsdruck nimmt zu",
            "Bildung beeinflusst Zukunft",
            "Wohlbefinden gehört zu erfolgreicher Bildung",
        ],
        size=25,
    )
    add_card(slide, 0.78, 5.25, 11.8, 1.05, fill=C.BLUE)
    text_box(slide, 1.08, 5.49, 11.2, 0.55, "Inwiefern beeinflussen strukturelle Unterschiede im deutschen Schulsystem das Wohlbefinden von Schülerinnen und Schülern?", size=22, bold=True, align=PP_ALIGN.CENTER, animate=True)
    add_notes(slide, notes)


def slide_3(prs):
    notes = (
        "Damit die Leitfrage verständlich wird, klärt die Facharbeit zunächst drei Grundlagen. "
        "Mit Strukturmerkmalen sind Rahmenbedingungen gemeint, zum Beispiel Schulformen, das "
        "Notensystem, Versetzungen und Abschlüsse. Wohlbefinden meint nicht nur gute Laune, sondern "
        "psychische Gesundheit, soziale Beziehungen und Lebenszufriedenheit im schulischen Kontext. "
        "Bildungsgerechtigkeit beschreibt faire Chancen unabhängig von der sozialen Herkunft. Diese "
        "drei Bereiche hängen zusammen: Strukturen können Bildungswege öffnen oder begrenzen, und "
        "sie können beeinflussen, ob Schülerinnen und Schüler Schule als unterstützend oder belastend "
        "erleben. Im nächsten Schritt schaue ich deshalb auf das deutsche Schulsystem selbst."
    )
    slide = begin_slide(prs, "Grundlagen", "Move In", C.YELLOW, notes)
    title(slide, "Grundlagen", subtitle="Drei Begriffe verbinden Struktur und Glück")
    center = add_card(slide, 5.15, 2.85, 2.15, 0.8, fill=C.WHITE)
    text_box(slide, 5.38, 3.08, 1.68, 0.32, "Schule", size=28, bold=True, align=PP_ALIGN.CENTER, animate=True)
    items = [
        ("Strukturmerkmale", ["Schulformen", "Notensystem", "Versetzungen", "Abschlüsse"], 0.8, 2.0, C.BLUE),
        ("Wohlbefinden", ["psychische Gesundheit", "soziale Beziehungen", "Lebenszufriedenheit"], 8.75, 1.85, C.PINK),
        ("Bildungsgerechtigkeit", ["faire Chancen", "unabhängig von sozialer Herkunft"], 4.0, 5.15, C.YELLOW),
    ]
    for head, bullets, x, y, fill in items:
        add_card(slide, x, y, 3.6, 1.55, fill=fill)
        text_box(slide, x + 0.25, y + 0.2, 3.1, 0.28, head, size=21, bold=True, align=PP_ALIGN.CENTER, animate=True)
        text_box(slide, x + 0.38, y + 0.67, 2.9, 0.55, " • ".join(bullets), size=15, color=C.DARK, align=PP_ALIGN.CENTER, animate=True)
        line(slide, 6.22, 3.25, x + 1.8, y + 0.78, C.MID, 1.2)
    add_notes(slide, notes)


def slide_4(prs):
    notes = (
        "Das deutsche Schulsystem ist durch Mehrgliedrigkeit geprägt. Die Facharbeit beschreibt, "
        "dass Schülerinnen und Schüler nach der Grundschule in verschiedene Schulformen wechseln. "
        "Das Gymnasium zielt auf das Abitur und akademische Bildungswege, ist aber auch mit hohen "
        "Leistungsanforderungen verbunden. Die Realschule verbindet allgemeine Bildung mit stärkerem "
        "Praxisbezug. Die Hauptschule war traditionell auf berufliche Wege ausgerichtet, wird aber "
        "in der Arbeit auch mit Stigmatisierung und veränderten Strukturen in den Bundesländern "
        "verbunden. Gesamtschulen ermöglichen längeres gemeinsames Lernen, und berufliche Schulen "
        "sind wichtig für das duale Ausbildungssystem. Auf dieser Grundlage wird verständlich, wo "
        "Leistungsdruck und Selektion im Alltag entstehen."
    )
    slide = begin_slide(prs, "Das deutsche Schulsystem", "Magic Move", C.BLUE, notes)
    title(slide, "Das deutsche Schulsystem", subtitle="Fünf Bildungswege im Überblick")
    cards = [
        ("Gymnasium", "Abitur\nakademische Wege\nhohe Anforderungen", C.BLUE, "🎓"),
        ("Realschule", "Mittlerer Abschluss\nmehr Praxisbezug\nBrücke zu Ausbildung", C.PINK, "🧭"),
        ("Hauptschule", "berufliche Ausrichtung\nAusbildung\nteils Stigmatisierung", C.YELLOW, "🛠"),
        ("Gesamtschule", "längeres gemeinsames Lernen\nKurse und Förderung\nmehrere Abschlüsse", C.LAV, "🤝"),
        ("Berufliche Schule", "duale Ausbildung\nBetrieb und Schule\nÜbergang in Arbeit", C.GREEN, "🏢"),
    ]
    x0 = 0.56
    for idx, (head, body, fill, icon) in enumerate(cards):
        x = x0 + idx * 2.55
        add_card(slide, x, 2.05, 2.25, 3.65, fill=fill)
        text_box(slide, x + 0.18, 2.32, 1.9, 0.34, icon, size=25, align=PP_ALIGN.CENTER, animate=True)
        text_box(slide, x + 0.18, 2.85, 1.9, 0.32, head, size=20, bold=True, align=PP_ALIGN.CENTER, animate=True)
        text_box(slide, x + 0.18, 3.52, 1.9, 1.2, body, size=14, align=PP_ALIGN.CENTER, animate=True)
    add_notes(slide, notes)


def slide_5(prs):
    notes = (
        "Ein zentraler Zusammenhang der Facharbeit ist der zwischen Leistungsdruck und Wohlbefinden. "
        "Leistungsdruck entsteht vor allem durch Noten, Prüfungen, Konkurrenz und Erwartungen von "
        "Eltern, Lehrkräften oder der Gesellschaft. Die Arbeit betont, dass Leistungsanforderungen "
        "nicht grundsätzlich negativ sind: Sie können Orientierung geben und Motivation fördern. "
        "Problematisch wird es aber, wenn Anforderungen als dauerhaft belastend erlebt werden. Dann "
        "können Stress, Schlafprobleme, Unsicherheit und ein geringeres Selbstwertgefühl entstehen. "
        "Besonders wichtig ist dabei die subjektive Wahrnehmung: dieselbe Anforderung kann für manche "
        "motivierend und für andere überfordernd sein. Danach geht es um den Bereich, der Belastungen "
        "abfedern kann: das Schulklima."
    )
    slide = begin_slide(prs, "Leistungsdruck und Wohlbefinden", "Scale", C.PINK, notes)
    title(slide, "Leistungsdruck und Wohlbefinden", subtitle="Bewertung kann antreiben und belasten")
    draw_stressed_student(slide, 9.1, 1.72)
    add_card(slide, 0.85, 2.0, 3.55, 2.65, fill=C.BLUE)
    text_box(slide, 1.15, 2.25, 2.9, 0.32, "Leistungsdruck", size=24, bold=True, align=PP_ALIGN.CENTER, animate=True)
    bullet_list(slide, 1.25, 2.85, 2.5, 1.25, ["Noten", "Prüfungen", "Konkurrenz", "Erwartungen"], size=20)
    add_card(slide, 4.95, 2.0, 3.55, 2.65, fill=C.YELLOW)
    text_box(slide, 5.25, 2.25, 2.9, 0.32, "Folgen", size=24, bold=True, align=PP_ALIGN.CENTER, animate=True)
    bullet_list(slide, 5.35, 2.85, 2.5, 1.25, ["Stress", "Schlafprobleme", "Unsicherheit", "geringeres Selbstwertgefühl"], size=20)
    add_card(slide, 1.05, 5.45, 7.75, 0.82, fill=C.PINK)
    text_box(slide, 1.32, 5.67, 7.2, 0.32, "Leistungsdruck kann Motivation fördern, aber auch das Wohlbefinden beeinträchtigen.", size=19, bold=True, align=PP_ALIGN.CENTER, animate=True)
    add_notes(slide, notes)


def slide_6(prs):
    notes = (
        "Neben Leistung spielt das Schulklima eine entscheidende Rolle. Die Facharbeit beschreibt "
        "Schulklima als Qualität der Beziehungen, Atmosphäre, Sicherheit und Zugehörigkeit. Ein "
        "positives Lehrer-Schüler-Verhältnis kann Stress reduzieren, weil Schülerinnen und Schüler "
        "sich ernst genommen und unterstützt fühlen. Mitbestimmung stärkt das Gefühl, nicht nur "
        "Objekt von Entscheidungen zu sein, sondern die eigene Schule mitgestalten zu können. "
        "Feedbackkultur bedeutet, dass Rückmeldungen nicht nur aus Noten bestehen, sondern auch "
        "Lernfortschritte und Stärken sichtbar machen. So wird Lernen stärker als Entwicklung "
        "verstanden. Nach diesen inhaltlichen Punkten beziehe ich kurz das Publikum ein."
    )
    slide = begin_slide(prs, "Schulklima", "Push", C.YELLOW, notes)
    title(slide, "Schulklima", subtitle="Beziehungen entscheiden mit über Wohlbefinden")
    draw_dialogue(slide, 9.35, 1.12)
    cards = [
        ("Lehrer-Schüler-Verhältnis", ["ernst genommen werden", "Unterstützung erhalten"], C.BLUE),
        ("Mitbestimmung", ["Meinung zählt", "Selbstwirksamkeit stärken"], C.PINK),
        ("Feedbackkultur", ["Lernfortschritte sehen", "Stärken entwickeln"], C.YELLOW),
    ]
    for idx, (head, points, fill) in enumerate(cards):
        x = 0.78 + idx * 4.0
        add_card(slide, x, 3.45, 3.45, 2.15, fill=fill)
        text_box(slide, x + 0.24, 3.74, 3.0, 0.45, head, size=21, bold=True, align=PP_ALIGN.CENTER, animate=True)
        chip(slide, x + 0.43, 4.42, 2.55, points[0], C.WHITE, size=14)
        chip(slide, x + 0.43, 4.92, 2.55, points[1], C.WHITE, size=14)
    add_notes(slide, notes)


def slide_7(prs):
    notes = (
        "An dieser Stelle eignet sich eine kurze Umfrage, weil das Thema nicht abstrakt bleiben soll. "
        "Ich frage: Was beeinflusst euer Wohlbefinden in der Schule am stärksten? Die vier Antworten "
        "greifen zentrale Bereiche der Facharbeit auf: Notendruck steht für Leistungsbewertung und "
        "Prüfungsstress, Lehrkräfte für das Lehrer-Schüler-Verhältnis, Mitschüler für soziale "
        "Beziehungen und Zeitstress für die Belastung durch schulische Anforderungen. Nach dem "
        "Handzeichen kann man kurz sichtbar machen, dass es vermutlich unterschiedliche Antworten "
        "gibt. Genau das passt zur Facharbeit: Wohlbefinden entsteht nicht aus einem einzelnen Faktor, "
        "sondern aus dem Zusammenspiel von Struktur, Leistung und Schulklima. Damit öffnet sich der "
        "Blick zum internationalen Vergleich."
    )
    slide = begin_slide(prs, "Kurze Umfrage", "Dissolve", C.BLUE, notes)
    title(slide, "Kurze Umfrage", subtitle="Publikum einbinden")
    add_card(slide, 2.15, 1.92, 9.05, 1.0, fill=C.WHITE)
    text_box(slide, 2.45, 2.18, 8.45, 0.42, "Was beeinflusst euer Wohlbefinden in der Schule am stärksten?", size=25, bold=True, align=PP_ALIGN.CENTER, animate=True)
    buttons = [
        ("📚 Notendruck", C.BLUE),
        ("👩‍🏫 Lehrkräfte", C.PINK),
        ("👥 Mitschüler", C.YELLOW),
        ("🕒 Zeitstress", C.LAV),
    ]
    for idx, (txt, fill) in enumerate(buttons):
        x = 1.35 + (idx % 2) * 5.55
        y = 3.55 + (idx // 2) * 0.92
        add_card(slide, x, y, 4.85, 0.62, fill=fill)
        text_box(slide, x + 0.18, y + 0.16, 4.5, 0.25, txt, size=19, bold=True, align=PP_ALIGN.CENTER, animate=True)
    text_box(slide, 5.35, 5.58, 2.6, 0.34, "Kurzes Handzeichen", size=18, bold=True, color=C.MID, align=PP_ALIGN.CENTER, animate=True)
    add_card(slide, 2.78, 6.18, 7.75, 0.54, fill=C.GREEN)
    text_box(slide, 3.05, 6.33, 7.2, 0.23, "Alle Faktoren spielen laut Facharbeit eine wichtige Rolle.", size=17, bold=True, align=PP_ALIGN.CENTER, animate=True)
    add_notes(slide, notes)


def slide_8(prs):
    notes = (
        "Der internationale Vergleich hilft, das deutsche System einzuordnen. Die Facharbeit nennt "
        "Deutschland als Beispiel für frühe Aufteilung und eine starke Rolle von Noten und Prüfungen. "
        "Finnland steht für längeres gemeinsames Lernen und eine spätere Leistungsdifferenzierung. "
        "Außerdem werden dort individuelle Lernfortschritte und qualitative Rückmeldungen stärker "
        "betont, besonders in frühen Schuljahren. Angelsächsische Systeme werden in der Arbeit als "
        "stärker durch Gesamtschulen, Wahlmöglichkeiten, Projektorientierung und praktische "
        "Anwendungen geprägt beschrieben. Wichtig ist aber: Die Facharbeit behauptet nicht, dass ein "
        "System perfekt ist. Jedes Bildungssystem hat eigene Herausforderungen. Der Vergleich zeigt "
        "vor allem, dass Strukturentscheidungen unterschiedlich gestaltet werden können."
    )
    slide = begin_slide(prs, "Internationaler Vergleich", "Pop", C.PINK, notes)
    title(slide, "Internationaler Vergleich", subtitle="Andere Strukturen, andere Schwerpunkte")
    cards = [
        ("Deutschland", ["frühe Aufteilung", "viele Noten"], C.BLUE, "DE"),
        ("Finnland", ["längeres gemeinsames Lernen", "spätere Leistungsdifferenzierung"], C.YELLOW, "FI"),
        ("Angelsächsische Systeme", ["mehr Praxis", "mehr Wahlmöglichkeiten"], C.PINK, "UK/US"),
    ]
    for idx, (head, bullets, fill, label) in enumerate(cards):
        x = 0.9 + idx * 4.1
        add_card(slide, x, 2.1, 3.55, 3.35, fill=fill)
        text_box(slide, x + 0.25, 2.42, 3.05, 0.34, label, size=23, bold=True, align=PP_ALIGN.CENTER, animate=True)
        text_box(slide, x + 0.25, 2.95, 3.05, 0.35, head, size=21, bold=True, align=PP_ALIGN.CENTER, animate=True)
        bullet_list(slide, x + 0.55, 3.65, 2.35, 0.95, bullets, size=18)
    text_box(slide, 5.0, 6.05, 3.3, 0.36, "kein System ist perfekt", size=18, bold=True, color=C.MID, align=PP_ALIGN.CENTER, animate=True)
    add_notes(slide, notes)


def slide_9(prs):
    notes = (
        "Aus der Analyse ergibt sich eine ausgewogene Bewertung. Zu den Stärken des deutschen "
        "Bildungssystems zählt die Facharbeit verschiedene Bildungswege. Diese können theoretisch "
        "passende Lernumgebungen und unterschiedliche Perspektiven ermöglichen. Außerdem wird das "
        "duale Ausbildungssystem hervorgehoben, weil es schulisches Lernen mit praktischer Ausbildung "
        "in Betrieben verbindet. Auch hohe akademische Standards, besonders im Gymnasium, werden als "
        "Stärke genannt. Gleichzeitig gibt es Schwächen: die frühe Selektion nach der Grundschule, "
        "Leistungsdruck durch Noten und Prüfungen sowie soziale Ungleichheiten. Diese Schwächen sind "
        "für die Leitfrage wichtig, weil sie direkt mit Wohlbefinden und Bildungsgerechtigkeit "
        "zusammenhängen. Daraus ergeben sich mögliche Reformansätze."
    )
    slide = begin_slide(prs, "Stärken und Schwächen", "Slide", C.YELLOW, notes)
    title(slide, "Stärken und Schwächen", subtitle="Eine kritische Bewertung")
    add_card(slide, 0.95, 2.0, 5.25, 3.9, fill=C.GREEN)
    text_box(slide, 1.28, 2.38, 4.5, 0.34, "Stärken", size=29, bold=True, align=PP_ALIGN.CENTER, animate=True)
    bullet_list(slide, 1.52, 3.12, 3.9, 1.55, ["verschiedene Bildungswege", "duales Ausbildungssystem", "hohe akademische Standards"], size=22)
    add_card(slide, 7.1, 2.0, 5.25, 3.9, fill=C.PINK)
    text_box(slide, 7.42, 2.38, 4.5, 0.34, "Schwächen", size=29, bold=True, align=PP_ALIGN.CENTER, animate=True)
    bullet_list(slide, 7.65, 3.12, 3.9, 1.55, ["frühe Selektion", "Leistungsdruck", "soziale Ungleichheiten"], size=22)
    add_notes(slide, notes)


def slide_10(prs):
    notes = (
        "Die Facharbeit diskutiert drei Reformansätze. Erstens könnten mehr Lebenskompetenzen in den "
        "Unterricht integriert werden. Dazu gehören zum Beispiel finanzielle Grundbildung, "
        "Medienkompetenz und praktische Alltagsfähigkeiten, also Inhalte mit stärkerem Bezug zum "
        "späteren Leben. Zweitens wird eine spätere Leistungsdifferenzierung genannt. Wenn Schülerinnen "
        "und Schüler länger gemeinsam lernen, hätten sie mehr Zeit, Fähigkeiten zu entwickeln und "
        "fundiertere Bildungsentscheidungen zu treffen. Drittens geht es um alternative Bewertungsformen: "
        "ausführliches Lernfeedback, Portfolioarbeit oder projektbasierte Bewertung können Lernprozesse "
        "stärker sichtbar machen. Diese Ansätze lösen nicht alle Probleme automatisch, zeigen aber, wie "
        "Bildung und Wohlbefinden stärker zusammengedacht werden könnten."
    )
    slide = begin_slide(prs, "Reformansätze", "Rotate", C.BLUE, notes)
    title(slide, "Reformansätze", subtitle="Wie Bildung und Wohlbefinden stärker zusammenfinden könnten")
    draw_future_arrow(slide, 4.55, 1.65)
    circles = [
        ("Mehr\nLebenskompetenzen", "Finanzbildung,\nMedienkompetenz,\nAlltagsfähigkeiten", 1.05, 3.45, C.YELLOW),
        ("Spätere\nLeistungsdifferenzierung", "mehr Zeit für\nEntwicklung und\nEntscheidungen", 5.05, 3.45, C.PINK),
        ("Alternative\nBewertung", "Feedback,\nPortfolio,\nProjekte", 9.05, 3.45, C.BLUE),
    ]
    for head, body, x, y, fill in circles:
        add_shadow(slide, x, y, 2.65, 2.65, MSO_SHAPE.OVAL)
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, i(x), i(y), i(2.65), i(2.65))
        circ.fill.solid()
        circ.fill.fore_color.rgb = fill
        circ.line.fill.background()
        text_box(slide, x + 0.22, y + 0.48, 2.2, 0.62, head, size=18, bold=True, align=PP_ALIGN.CENTER, animate=True)
        text_box(slide, x + 0.25, y + 1.34, 2.15, 0.72, body, size=14, align=PP_ALIGN.CENTER, animate=True)
    add_notes(slide, notes)


def slide_11(prs):
    notes = (
        "Zum Schluss lässt sich die Leitfrage so beantworten: Schule beeinflusst mehr als nur Noten. "
        "Die Facharbeit zeigt, dass Strukturen wie Schulformen, frühe Differenzierung, Notensystem und "
        "Übergangsentscheidungen auf Lernbedingungen und Wohlbefinden wirken können. Leistungsdruck "
        "und Schulklima spielen dabei eine zentrale Rolle. Noten und Prüfungen können Orientierung "
        "geben, aber auch Stress und Unsicherheit verstärken. Gleichzeitig können unterstützende "
        "Beziehungen, Mitbestimmung und gute Feedbackkultur Wohlbefinden fördern. Reformen könnten "
        "daher versuchen, Bildung und Glück stärker miteinander zu verbinden: nicht als Gegensatz zu "
        "Leistung, sondern als Voraussetzung dafür, dass Lernen langfristig gelingen kann. Damit danke "
        "ich euch für eure Aufmerksamkeit."
    )
    slide = begin_slide(prs, "Fazit", "Fade", C.PINK, notes)
    text_box(slide, 0.9, 0.88, 8.65, 0.7, "Schule beeinflusst mehr als nur Noten.", size=42, bold=True, font=FONT_HEAD, animate=True)
    add_card(slide, 0.95, 2.12, 7.45, 3.1, fill=C.WHITE)
    statements = [
        "✓ Strukturen wirken auf das Wohlbefinden.",
        "✓ Leistungsdruck und Schulklima spielen eine zentrale Rolle.",
        "✓ Reformen könnten Bildung und Glück stärker miteinander verbinden.",
    ]
    for idx, stmt in enumerate(statements):
        text_box(slide, 1.32, 2.62 + idx * 0.78, 6.65, 0.36, stmt, size=22, bold=True, animate=True)
    # Calm pastel illustration: open path, sun, and three connected learners.
    sun = slide.shapes.add_shape(MSO_SHAPE.OVAL, i(10.45), i(1.05), i(1.15), i(1.15))
    sun.fill.solid()
    sun.fill.fore_color.rgb = C.YELLOW
    sun.line.fill.background()
    line(slide, 8.95, 5.25, 12.45, 5.25, C.MID, 2.2)
    for x, color in [(9.25, C.BLUE), (10.55, C.PINK), (11.85, C.YELLOW)]:
        draw_person(slide, x, 3.6, 1.15, color, "smile")
    text_box(slide, 0.98, 6.4, 5.4, 0.38, "Danke für eure Aufmerksamkeit!", size=25, bold=True, color=C.MID, animate=True)
    add_notes(slide, notes)


TRANSITION_XML = {
    "Fade": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:fade/></p:transition>',
    "Zoom": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:zoom dir="in"/></p:transition>',
    "Move In": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:cover dir="l"/></p:transition>',
    "Magic Move": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" spd="med"><p159:morph option="byObject"/></p:transition>',
    "Scale": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:zoom dir="out"/></p:transition>',
    "Push": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:push dir="l"/></p:transition>',
    "Dissolve": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:dissolve/></p:transition>',
    "Pop": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="fast"><p:plus/></p:transition>',
    "Slide": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:wipe dir="l"/></p:transition>',
    "Rotate": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:newsflash/></p:transition>',
}


def make_timing(shape_ids: list[int]) -> etree._Element:
    ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    P = f"{{{ns}}}"

    def e(name, **attrs):
        return etree.Element(P + name, **{k: str(v) for k, v in attrs.items()})

    timing = e("timing")
    tn_lst = etree.SubElement(timing, P + "tnLst")
    par_root = etree.SubElement(tn_lst, P + "par")
    ctn_root = etree.SubElement(par_root, P + "cTn", id="1", dur="indefinite", restart="never", nodeType="tmRoot")
    child_root = etree.SubElement(ctn_root, P + "childTnLst")
    seq = etree.SubElement(child_root, P + "seq", concurrent="1", nextAc="seek")
    ctn_seq = etree.SubElement(seq, P + "cTn", id="2", dur="indefinite", nodeType="mainSeq")
    child_seq = etree.SubElement(ctn_seq, P + "childTnLst")

    next_id = 3
    for idx, spid in enumerate(shape_ids):
        par = etree.SubElement(child_seq, P + "par")
        ctn = etree.SubElement(
            par,
            P + "cTn",
            id=str(next_id),
            fill="hold",
            presetID="10",
            presetClass="entr",
            presetSubtype="0",
            nodeType="clickEffect" if idx == 0 else "afterEffect",
        )
        next_id += 1
        st = etree.SubElement(ctn, P + "stCondLst")
        etree.SubElement(st, P + "cond", delay="0" if idx == 0 else "250")
        child = etree.SubElement(ctn, P + "childTnLst")
        anim = etree.SubElement(child, P + "animEffect", transition="in", filter="fade")
        bhvr = etree.SubElement(anim, P + "cBhvr")
        etree.SubElement(bhvr, P + "cTn", id=str(next_id), dur="450", fill="hold")
        next_id += 1
        tgt = etree.SubElement(bhvr, P + "tgtEl")
        etree.SubElement(tgt, P + "spTgt", spid=str(spid))

    bld_lst = etree.SubElement(timing, P + "bldLst")
    for spid in shape_ids:
        etree.SubElement(bld_lst, P + "bldP", spid=str(spid), grpId="0", build="allAtOnce")
    return timing


def patch_pptx(path: Path):
    tmp = path.with_suffix(".patched.pptx")
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("ppt/slides/slide") and item.filename.endswith(".xml"):
                slide_no = int(Path(item.filename).stem.replace("slide", ""))
                info = slide_infos[slide_no - 1]
                root = etree.fromstring(data)
                for old in root.xpath("./p:transition | ./p:timing", namespaces=ns):
                    root.remove(old)
                transition = etree.fromstring(TRANSITION_XML[info.transition].encode("utf-8"))
                c_sld = root.find("p:cSld", namespaces=ns)
                clr = root.find("p:clrMapOvr", namespaces=ns)
                insert_idx = list(root).index(clr) + 1 if clr is not None else list(root).index(c_sld) + 1
                root.insert(insert_idx, transition)
                if info.animation_ids:
                    root.insert(insert_idx + 1, make_timing(info.animation_ids))
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    tmp.replace(path)


def write_notes_md():
    lines = [
        "# Sprechnotizen",
        "",
        "Die folgenden Moderationsnotizen sind frei vortragbar und auf ca. eine Minute pro Folie angelegt.",
        "",
    ]
    for idx, info in enumerate(slide_infos, start=1):
        lines.extend(
            [
                f"## Folie {idx}: {info.title}",
                "",
                f"**Übergang:** {info.transition}",
                "",
                info.notes,
                "",
            ]
        )
    NOTES_PATH.write_text("\n".join(lines), encoding="utf-8")


def build():
    OUT_DIR.mkdir(exist_ok=True)
    prs = Presentation()
    prs.slide_width = i(SLIDE_W)
    prs.slide_height = i(SLIDE_H)

    for fn in (slide_1, slide_2, slide_3, slide_4, slide_5, slide_6, slide_7, slide_8, slide_9, slide_10, slide_11):
        fn(prs)

    raw = OUT_DIR / "raw.pptx"
    prs.save(raw)
    shutil.copy(raw, PPTX_PATH)
    raw.unlink()
    patch_pptx(PPTX_PATH)
    write_notes_md()


if __name__ == "__main__":
    build()
