from __future__ import annotations

from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Inches, Pt


# Pastellpalette laut Vorgabe
COLOR_BG_BLUE = RGBColor(0xCF, 0xEF, 0xFF)
COLOR_BG_PINK = RGBColor(0xFF, 0xD9, 0xE8)
COLOR_BG_YELLOW = RGBColor(0xFF, 0xF7, 0xC7)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_TEXT = RGBColor(0x33, 0x33, 0x33)
COLOR_ACCENT_GREEN = RGBColor(0xD8, 0xF1, 0xD2)


def add_round_rect(slide, left, top, width, height, color, line_color=COLOR_WHITE, line_width=1.5):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = line_color
    shape.line.width = Pt(line_width)
    return shape


def add_title(slide, text: str, top=Inches(0.35), size=44):
    box = slide.shapes.add_textbox(Inches(0.7), top, Inches(11.8), Inches(1.0))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Avenir Next"
    p.font.bold = True
    p.font.size = Pt(size)
    p.font.color.rgb = COLOR_TEXT
    return box


def add_subtitle(slide, text: str, top=Inches(1.25), size=22):
    box = slide.shapes.add_textbox(Inches(0.7), top, Inches(8.2), Inches(1.0))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Avenir Next"
    p.font.size = Pt(size)
    p.font.color.rgb = COLOR_TEXT
    return box


def add_bullets(
    slide,
    title: str,
    bullets: Iterable[str],
    left,
    top,
    width,
    height,
    title_size=24,
    bullet_size=21,
):
    panel = add_round_rect(slide, left, top, width, height, COLOR_WHITE)
    tf = panel.text_frame
    tf.clear()

    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.name = "Avenir Next"
    p0.font.bold = True
    p0.font.size = Pt(title_size)
    p0.font.color.rgb = COLOR_TEXT

    for item in bullets:
        p = tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Avenir Next"
        p.font.size = Pt(bullet_size)
        p.font.color.rgb = COLOR_TEXT
    return panel


def decorate_background(slide, variant=0):
    if variant % 3 == 0:
        add_round_rect(slide, Inches(10.9), Inches(0.2), Inches(2.1), Inches(2.1), COLOR_BG_BLUE)
        add_round_rect(slide, Inches(0.2), Inches(5.8), Inches(2.5), Inches(1.3), COLOR_BG_PINK)
    elif variant % 3 == 1:
        add_round_rect(slide, Inches(10.6), Inches(5.9), Inches(2.4), Inches(1.2), COLOR_BG_YELLOW)
        add_round_rect(slide, Inches(0.25), Inches(0.2), Inches(2.8), Inches(1.2), COLOR_BG_BLUE)
    else:
        add_round_rect(slide, Inches(11.0), Inches(5.8), Inches(1.9), Inches(1.1), COLOR_BG_PINK)
        add_round_rect(slide, Inches(0.2), Inches(0.15), Inches(2.0), Inches(1.1), COLOR_BG_YELLOW)


def add_visual_school(slide):
    card = add_round_rect(slide, Inches(8.8), Inches(1.5), Inches(4.0), Inches(4.8), COLOR_BG_BLUE)
    bld = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(9.4), Inches(2.7), Inches(2.8), Inches(2.2))
    bld.fill.solid()
    bld.fill.fore_color.rgb = COLOR_WHITE
    bld.line.color.rgb = COLOR_TEXT
    roof = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE, Inches(9.3), Inches(2.0), Inches(3.0), Inches(1.0))
    roof.fill.solid()
    roof.fill.fore_color.rgb = COLOR_BG_PINK
    roof.line.color.rgb = COLOR_TEXT
    for r in range(2):
        for c in range(3):
            win = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                Inches(9.65 + 0.75 * c),
                Inches(3.0 + 0.65 * r),
                Inches(0.35),
                Inches(0.35),
            )
            win.fill.solid()
            win.fill.fore_color.rgb = COLOR_BG_YELLOW
            win.line.color.rgb = COLOR_TEXT
    card.text_frame.text = ""


def add_transition(slide, transition_type: str):
    for child in list(slide._element):
        if child.tag.endswith("transition"):
            slide._element.remove(child)
    trans_xml = (
        f'<p:transition {nsdecls("p")} advClick="1">'
        f"<p:{transition_type}/>"
        f"</p:transition>"
    )
    trans = parse_xml(trans_xml)
    slide._element.insert(len(slide._element), trans)


def add_notes(slide, text: str):
    notes = slide.notes_slide.notes_text_frame
    notes.clear()
    notes.text = text


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    layout = prs.slide_layouts[6]

    # Slide 1
    s1 = prs.slides.add_slide(layout)
    decorate_background(s1, 0)
    add_title(s1, "Wie Schule wirkt", top=Inches(0.8), size=56)
    add_subtitle(
        s1,
        "Strukturelle Unterschiede und ihr Effekt auf das Wohlbefinden\nvon Schülerinnen und Schülern",
        top=Inches(2.0),
        size=24,
    )
    add_subtitle(s1, "Seminarfachkurs Glück", top=Inches(3.35), size=24)
    add_visual_school(s1)
    add_transition(s1, "fade")
    add_notes(
        s1,
        "Willkommen zu meiner Präsentation im Seminarfachkurs Glück. "
        "Ich untersuche heute, wie Schule auf das Wohlbefinden von Schülerinnen und Schülern wirkt. "
        "Im Mittelpunkt steht die Frage, ob strukturelle Unterschiede im deutschen Schulsystem nicht nur "
        "Bildungswege, sondern auch psychische und soziale Erfahrungen prägen. "
        "Die Facharbeit zeigt: Schule ist mehr als Unterricht und Noten. Sie ist ein zentraler Lebensraum, "
        "in dem Selbstwert, Motivation und Zukunftsperspektiven mitgeformt werden. "
        "Genau deshalb ist der Blick auf Strukturfragen so wichtig. "
        "Als nächstes zeige ich, warum dieses Thema gesellschaftlich relevant ist und welche Leitfrage die Arbeit leitet.",
    )

    # Slide 2
    s2 = prs.slides.add_slide(layout)
    decorate_background(s2, 1)
    add_title(s2, "Warum ist das Thema wichtig?", top=Inches(0.55), size=44)
    add_bullets(
        s2,
        "Relevanz",
        [
            "Schule prägt den Alltag junger Menschen.",
            "Leistungsdruck wird von vielen Jugendlichen als Belastung erlebt.",
            "Bildung beeinflusst soziale Teilhabe und Zukunftschancen.",
            "Wohlbefinden ist Teil erfolgreicher Bildung.",
        ],
        Inches(0.8),
        Inches(1.55),
        Inches(7.0),
        Inches(3.8),
    )
    q = add_round_rect(s2, Inches(0.8), Inches(5.55), Inches(11.8), Inches(1.15), COLOR_BG_YELLOW, line_color=COLOR_TEXT)
    q.text_frame.text = (
        "Inwiefern beeinflussen strukturelle Unterschiede im deutschen Schulsystem "
        "das Wohlbefinden von Schülerinnen und Schülern?"
    )
    qtf = q.text_frame.paragraphs[0]
    qtf.font.name = "Avenir Next"
    qtf.font.size = Pt(24)
    qtf.font.bold = True
    qtf.font.color.rgb = COLOR_TEXT
    qtf.alignment = PP_ALIGN.CENTER

    scale = add_round_rect(s2, Inches(8.25), Inches(1.55), Inches(4.35), Inches(3.8), COLOR_BG_PINK)
    scale.text_frame.text = "⚖ Bildung\nund\nWohlbefinden"
    for p in scale.text_frame.paragraphs:
        p.font.name = "Avenir Next"
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT
        p.alignment = PP_ALIGN.CENTER
    add_transition(s2, "zoom")
    add_notes(
        s2,
        "Die Facharbeit zeigt deutlich, warum dieses Thema relevant ist. "
        "Schule ist ein täglicher Lebensraum und wirkt damit auf Lernprozesse, Beziehungen und psychische Gesundheit. "
        "Gleichzeitig steht sie unter dem Anspruch, auf Zukunft und Beruf vorzubereiten. "
        "Genau in diesem Spannungsfeld zwischen Bildungsauftrag und Lebensrealität entsteht die zentrale Problematik: "
        "Leistungsdruck, Vergleich und Selektion können belasten, obwohl Bildung eigentlich fördern soll. "
        "Deshalb verbindet die Arbeit individuelle Perspektiven mit gesellschaftlichen Fragen der Bildungsgerechtigkeit. "
        "Die hervorgehobene Leitfrage ist der rote Faden der gesamten Untersuchung. "
        "Im nächsten Schritt kläre ich die grundlegenden Begriffe, mit denen die Arbeit operiert.",
    )

    # Slide 3
    s3 = prs.slides.add_slide(layout)
    decorate_background(s3, 2)
    add_title(s3, "Grundlagen", top=Inches(0.55), size=44)

    c1 = add_round_rect(s3, Inches(0.8), Inches(1.6), Inches(4.0), Inches(4.8), COLOR_BG_BLUE)
    c1.text_frame.text = (
        "Strukturmerkmale\n"
        "• Schulformen\n"
        "• Notensystem\n"
        "• Versetzungen\n"
        "• Abschlüsse"
    )
    c2 = add_round_rect(s3, Inches(4.95), Inches(1.6), Inches(4.0), Inches(4.8), COLOR_BG_PINK)
    c2.text_frame.text = (
        "Wohlbefinden\n"
        "• psychische Gesundheit\n"
        "• soziale Beziehungen\n"
        "• Lebenszufriedenheit"
    )
    c3 = add_round_rect(s3, Inches(9.1), Inches(1.6), Inches(3.45), Inches(4.8), COLOR_BG_YELLOW)
    c3.text_frame.text = (
        "Bildungsgerechtigkeit\n"
        "• faire Chancen\n"
        "• unabhängig von\n  sozialer Herkunft"
    )
    for box in (c1, c2, c3):
        for i, p in enumerate(box.text_frame.paragraphs):
            p.font.name = "Avenir Next"
            p.font.color.rgb = COLOR_TEXT
            p.font.bold = i == 0
            p.font.size = Pt(25 if i == 0 else 20)
            p.alignment = PP_ALIGN.LEFT
    add_transition(s3, "cover")
    add_notes(
        s3,
        "Bevor wir in die Analyse gehen, definiert die Facharbeit drei zentrale Begriffe. "
        "Erstens die Strukturmerkmale von Schule: also institutionelle Regeln wie Schulformen, Notensystem, "
        "Versetzungsordnung und Abschlüsse. Diese Elemente steuern Bildungswege und Selektionsprozesse. "
        "Zweitens das Wohlbefinden: Es umfasst nicht nur Leistungen, sondern auch psychische Gesundheit, "
        "soziale Einbindung und subjektive Lebenszufriedenheit. "
        "Und drittens Bildungsgerechtigkeit: die faire Verteilung von Chancen unabhängig von sozialer Herkunft. "
        "Gerade die Verbindung dieser drei Ebenen ist entscheidend, um die Leitfrage zu beantworten. "
        "Auf dieser Grundlage schauen wir uns als Nächstes die Schulformen im deutschen System an.",
    )

    # Slide 4
    s4 = prs.slides.add_slide(layout)
    decorate_background(s4, 0)
    add_title(s4, "Das deutsche Schulsystem", top=Inches(0.55), size=44)
    forms = [
        ("Gymnasium", "hohe akademische Anforderungen, Ziel: Abitur"),
        ("Realschule", "mittlerer Abschluss, Verbindung von Theorie und Praxis"),
        ("Hauptschule", "traditionell praxisnah, in vielen Ländern umstrukturiert"),
        ("Gesamtschule", "längeres gemeinsames Lernen mit Kurssystemen"),
        ("Berufliche Schule", "duales System aus Betrieb und Schule"),
    ]
    y = 1.5
    for i, (title, text) in enumerate(forms):
        card = add_round_rect(
            s4,
            Inches(0.9 + (i % 2) * 6.1),
            Inches(y + (i // 2) * 1.8),
            Inches(5.8),
            Inches(1.5),
            [COLOR_BG_BLUE, COLOR_BG_PINK, COLOR_BG_YELLOW, COLOR_WHITE, COLOR_ACCENT_GREEN][i],
            line_color=COLOR_TEXT,
        )
        tf = card.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Avenir Next"
        p0.font.bold = True
        p0.font.size = Pt(24)
        p0.font.color.rgb = COLOR_TEXT
        p1 = tf.add_paragraph()
        p1.text = text
        p1.font.name = "Avenir Next"
        p1.font.size = Pt(17)
        p1.font.color.rgb = COLOR_TEXT
    add_transition(s4, "push")
    add_notes(
        s4,
        "Die Facharbeit beschreibt das deutsche System als historisch gewachsen und mehrgliedrig. "
        "Gymnasium, Realschule, Hauptschule, Gesamtschule und berufliche Schulen verfolgen jeweils "
        "unterschiedliche Bildungsziele. Das kann gezielte Förderung ermöglichen, führt aber auch zu "
        "frühen Weichenstellungen. Besonders wichtig ist: Die Aufteilung erfolgt in vielen Bundesländern "
        "bereits nach der Grundschule. Genau diese frühe Differenzierung ist ein zentraler Streitpunkt in der "
        "bildungspolitischen Diskussion. Gleichzeitig zeigt das duale System eine oft positiv bewertete Stärke "
        "der deutschen Bildungsstruktur. "
        "Im nächsten Schritt fokussiere ich den Bereich, der in der Facharbeit als besonders belastend "
        "beschrieben wird: Leistungsdruck und seine Folgen für das Wohlbefinden.",
    )

    # Slide 5
    s5 = prs.slides.add_slide(layout)
    decorate_background(s5, 1)
    add_title(s5, "Leistungsdruck und Wohlbefinden", top=Inches(0.55), size=44)

    left = add_round_rect(s5, Inches(0.9), Inches(1.7), Inches(5.8), Inches(4.2), COLOR_BG_BLUE)
    left.text_frame.text = "Leistungsdruck\n• Noten\n• Prüfungen\n• Konkurrenz\n• Erwartungen"
    right = add_round_rect(s5, Inches(6.9), Inches(1.7), Inches(5.5), Inches(4.2), COLOR_BG_PINK)
    right.text_frame.text = "Folgen\n• Stress\n• Schlafprobleme\n• Unsicherheit\n• geringeres Selbstwertgefühl"
    for bx in (left, right):
        for i, p in enumerate(bx.text_frame.paragraphs):
            p.font.name = "Avenir Next"
            p.font.color.rgb = COLOR_TEXT
            p.font.bold = i == 0
            p.font.size = Pt(28 if i == 0 else 23)

    note = add_round_rect(s5, Inches(0.9), Inches(6.15), Inches(11.5), Inches(0.9), COLOR_BG_YELLOW, line_color=COLOR_TEXT)
    note.text_frame.text = "Leistungsdruck kann Motivation fördern, aber auch das Wohlbefinden beeinträchtigen."
    np = note.text_frame.paragraphs[0]
    np.font.name = "Avenir Next"
    np.font.size = Pt(22)
    np.font.bold = True
    np.font.color.rgb = COLOR_TEXT
    np.alignment = PP_ALIGN.CENTER
    add_transition(s5, "split")
    add_notes(
        s5,
        "Die Facharbeit zeigt klar: Leistungsdruck ist ein Schlüsselfaktor für schulisches Wohlbefinden. "
        "Noten, Prüfungen und Versetzungsentscheidungen sind im deutschen System zentral und haben konkrete Folgen "
        "für Bildungswege. Viele Jugendliche nehmen diese Anforderungen als Stressquelle wahr. "
        "Genannt werden unter anderem Unsicherheit, psychosomatische Beschwerden und ein sinkendes Selbstwertgefühl, "
        "vor allem bei wiederholten Misserfolgserfahrungen. Gleichzeitig ist wichtig, dass Leistungsanforderungen "
        "nicht nur negativ sind: Sie können auch motivieren. Entscheidend ist also das Ausmaß und die subjektive Wahrnehmung. "
        "Genau hier kommt der nächste Faktor ins Spiel: das Schulklima und die Qualität sozialer Beziehungen.",
    )

    # Slide 6
    s6 = prs.slides.add_slide(layout)
    decorate_background(s6, 2)
    add_title(s6, "Schulklima", top=Inches(0.55), size=44)

    cards = [
        (
            "Lehrer-Schüler-Verhältnis",
            ["respektvolle Beziehung stärkt Vertrauen", "Unterstützung reduziert Stress", "Ungerechtigkeit belastet deutlich"],
            COLOR_BG_BLUE,
        ),
        (
            "Mitbestimmung",
            ["Schülervertretung und Klassenrat", "Beteiligung stärkt Selbstwirksamkeit", "Meinung zählt im Schulalltag"],
            COLOR_BG_PINK,
        ),
        (
            "Feedbackkultur",
            ["nicht nur Noten, auch qualitative Rückmeldung", "formative Begleitung von Lernprozessen", "Stärken gezielt entwickeln"],
            COLOR_BG_YELLOW,
        ),
    ]
    for i, (t, items, col) in enumerate(cards):
        c = add_round_rect(s6, Inches(0.8 + i * 4.25), Inches(1.7), Inches(4.0), Inches(4.9), col)
        tf = c.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = t
        p0.font.name = "Avenir Next"
        p0.font.bold = True
        p0.font.size = Pt(24)
        p0.font.color.rgb = COLOR_TEXT
        for item in items:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.name = "Avenir Next"
            p.font.size = Pt(17)
            p.font.color.rgb = COLOR_TEXT
    add_transition(s6, "push")
    add_notes(
        s6,
        "Neben Leistungsdruck betont die Facharbeit das Schulklima als entscheidenden Einflussfaktor. "
        "Ein unterstützendes Lehrer-Schüler-Verhältnis wirkt stabilisierend, weil Schülerinnen und Schüler "
        "sich ernst genommen und begleitet fühlen. Zweitens spielt Mitbestimmung eine wichtige Rolle: "
        "Wenn Lernende beteiligt werden, steigt das Gefühl von Selbstwirksamkeit. "
        "Drittens ist die Feedbackkultur zentral. Qualitative Rückmeldungen ergänzen Noten und helfen, "
        "Lernfortschritte differenziert wahrzunehmen. Insgesamt zeigt sich: Wohlbefinden entsteht nicht nur durch "
        "individuelle Stärke, sondern durch die Qualität schulischer Beziehungen und Strukturen. "
        "Nach diesem inhaltlichen Teil binde ich jetzt kurz das Publikum mit einer Frage ein.",
    )

    # Slide 7
    s7 = prs.slides.add_slide(layout)
    decorate_background(s7, 0)
    add_title(s7, "Kurze Umfrage", top=Inches(0.7), size=48)

    qbox = add_round_rect(s7, Inches(1.1), Inches(1.75), Inches(11.1), Inches(1.25), COLOR_WHITE, line_color=COLOR_TEXT)
    qbox.text_frame.text = "Was beeinflusst euer Wohlbefinden in der Schule am stärksten?"
    qp = qbox.text_frame.paragraphs[0]
    qp.font.name = "Avenir Next"
    qp.font.size = Pt(28)
    qp.font.bold = True
    qp.font.color.rgb = COLOR_TEXT
    qp.alignment = PP_ALIGN.CENTER

    btns = [("📚 Notendruck", COLOR_BG_BLUE), ("👩‍🏫 Lehrkräfte", COLOR_BG_PINK), ("👥 Mitschüler", COLOR_BG_YELLOW), ("🕒 Zeitstress", COLOR_ACCENT_GREEN)]
    for i, (label, col) in enumerate(btns):
        b = add_round_rect(s7, Inches(1.4 + (i % 2) * 5.8), Inches(3.35 + (i // 2) * 1.35), Inches(4.8), Inches(1.0), col, line_color=COLOR_TEXT)
        b.text_frame.text = label
        bp = b.text_frame.paragraphs[0]
        bp.font.name = "Avenir Next"
        bp.font.size = Pt(24)
        bp.font.bold = True
        bp.font.color.rgb = COLOR_TEXT
        bp.alignment = PP_ALIGN.CENTER

    info = slide_textbox(s7, Inches(5.1), Inches(6.1), Inches(3.3), Inches(0.55), "Kurzes Handzeichen", 19, bold=False)
    info.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    result = slide_textbox(
        s7,
        Inches(2.3),
        Inches(6.6),
        Inches(8.8),
        Inches(0.6),
        "Alle Faktoren spielen laut Facharbeit eine wichtige Rolle.",
        21,
        bold=True,
    )
    result.fill.solid()
    result.fill.fore_color.rgb = COLOR_WHITE
    result.line.color.rgb = COLOR_TEXT
    add_transition(s7, "dissolve")
    add_notes(
        s7,
        "An dieser Stelle binde ich euch kurz ein. Die Frage lautet: Was beeinflusst euer Wohlbefinden in der Schule "
        "am stärksten? Wir machen dazu ein kurzes Handzeichen. "
        "Die vier Optionen greifen zentrale Faktoren aus der Facharbeit auf: Notendruck, Beziehung zu Lehrkräften, "
        "Mitschülerinnen und Mitschüler sowie Zeit- und Leistungsstress. "
        "Wichtig ist: Die Arbeit kommt nicht zu dem Ergebnis, dass nur ein Faktor entscheidend ist. "
        "Vielmehr wirken diese Aspekte zusammen und verstärken sich teilweise gegenseitig. "
        "Damit knüpfen wir direkt an den internationalen Vergleich an und schauen, wie andere Systeme mit "
        "Leistungsdruck und Strukturfragen umgehen.",
    )

    # Slide 8
    s8 = prs.slides.add_slide(layout)
    decorate_background(s8, 1)
    add_title(s8, "Internationaler Vergleich", top=Inches(0.55), size=44)
    data8 = [
        ("Deutschland", ["frühe Aufteilung", "viele Noten"], COLOR_BG_BLUE),
        ("Finnland", ["längeres gemeinsames Lernen", "spätere Leistungsdifferenzierung"], COLOR_BG_PINK),
        ("angelsächsische Systeme", ["mehr Praxis", "mehr Wahlmöglichkeiten"], COLOR_BG_YELLOW),
    ]
    for i, (t, items, col) in enumerate(data8):
        c = add_round_rect(s8, Inches(0.9 + i * 4.15), Inches(1.8), Inches(3.95), Inches(4.5), col, line_color=COLOR_TEXT)
        tf = c.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = t
        p0.font.name = "Avenir Next"
        p0.font.bold = True
        p0.font.size = Pt(24)
        p0.font.color.rgb = COLOR_TEXT
        for item in items:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.name = "Avenir Next"
            p.font.size = Pt(20)
            p.font.color.rgb = COLOR_TEXT
    bottom = add_round_rect(s8, Inches(3.55), Inches(6.45), Inches(6.2), Inches(0.8), COLOR_WHITE, line_color=COLOR_TEXT)
    bottom.text_frame.text = "Kein System ist perfekt."
    bp = bottom.text_frame.paragraphs[0]
    bp.font.name = "Avenir Next"
    bp.font.bold = True
    bp.font.size = Pt(24)
    bp.font.color.rgb = COLOR_TEXT
    bp.alignment = PP_ALIGN.CENTER
    add_transition(s8, "randomBar")
    add_notes(
        s8,
        "Der internationale Vergleich hilft, die deutsche Struktur besser einzuordnen. "
        "Deutschland ist durch frühe Differenzierung und starke Notenorientierung geprägt. "
        "Finnland setzt auf längeres gemeinsames Lernen und eine spätere Aufteilung in Bildungswege. "
        "In angelsächsischen Systemen fallen vor allem Praxisbezug und Wahlmöglichkeiten auf. "
        "Die Facharbeit zeigt dabei keine einfache Rangliste, sondern unterschiedliche Schwerpunkte. "
        "Entscheidend ist, welche Ziele ein System priorisiert: akademische Selektion, soziale Durchlässigkeit "
        "oder praxisnahe Kompetenzentwicklung. Deshalb lautet die zentrale Erkenntnis hier: "
        "Es gibt kein perfektes Modell, aber es gibt lernenswerte Ansätze. "
        "Darauf aufbauend fasse ich jetzt die Stärken und Schwächen des deutschen Systems zusammen.",
    )

    # Slide 9
    s9 = prs.slides.add_slide(layout)
    decorate_background(s9, 2)
    add_title(s9, "Stärken und Schwächen", top=Inches(0.55), size=44)
    st = add_round_rect(s9, Inches(0.9), Inches(1.7), Inches(5.8), Inches(4.8), COLOR_ACCENT_GREEN, line_color=COLOR_TEXT)
    st.text_frame.text = (
        "Stärken\n"
        "• verschiedene Bildungswege\n"
        "• duales Ausbildungssystem\n"
        "• hohe akademische Standards"
    )
    sw = add_round_rect(s9, Inches(6.95), Inches(1.7), Inches(5.4), Inches(4.8), COLOR_BG_PINK, line_color=COLOR_TEXT)
    sw.text_frame.text = (
        "Schwächen\n"
        "• frühe Selektion\n"
        "• Leistungsdruck\n"
        "• soziale Ungleichheiten"
    )
    for bx in (st, sw):
        for i, p in enumerate(bx.text_frame.paragraphs):
            p.font.name = "Avenir Next"
            p.font.color.rgb = COLOR_TEXT
            p.font.bold = i == 0
            p.font.size = Pt(31 if i == 0 else 24)
    add_transition(s9, "wipe")
    add_notes(
        s9,
        "Die kritische Bewertung der Facharbeit ist ausgewogen. "
        "Zu den Stärken zählen die differenzierten Bildungswege, das international anerkannte duale "
        "Ausbildungssystem und hohe akademische Standards. "
        "Gleichzeitig werden strukturelle Schwächen klar benannt: die frühe Selektion, anhaltender "
        "Leistungsdruck und die enge Kopplung von Bildungserfolg an soziale Herkunft. "
        "Damit wird deutlich: Das deutsche System ist weder grundsätzlich schlecht noch unproblematisch, "
        "sondern in sich widersprüchlich. Genau aus dieser Spannung ergeben sich Reformansätze, "
        "die Bildungsqualität und Wohlbefinden stärker zusammenführen sollen. "
        "Diese drei Reformideen zeige ich auf der nächsten Folie.",
    )

    # Slide 10
    s10 = prs.slides.add_slide(layout)
    decorate_background(s10, 0)
    add_title(s10, "Reformansätze", top=Inches(0.55), size=44)

    circles = [
        ("Mehr Lebens-\nkompetenzen", "z. B. finanzielle Grundbildung,\nMedienkompetenz, Alltagskompetenzen", COLOR_BG_BLUE),
        ("Spätere Leistungs-\ndifferenzierung", "längeres gemeinsames Lernen,\nmehr Zeit für Entwicklung", COLOR_BG_PINK),
        ("Alternative\nBewertungsformen", "mehr Lernfeedback,\nPortfolio- oder projektbasierte Bewertung", COLOR_BG_YELLOW),
    ]
    for i, (title, desc, col) in enumerate(circles):
        c = slide_circle(s10, Inches(0.9 + i * 4.15), Inches(1.8), Inches(3.8), col)
        tf = c.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Avenir Next"
        p0.font.bold = True
        p0.font.size = Pt(21)
        p0.font.color.rgb = COLOR_TEXT
        p0.alignment = PP_ALIGN.CENTER
        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.name = "Avenir Next"
        p1.font.size = Pt(14)
        p1.font.color.rgb = COLOR_TEXT
        p1.alignment = PP_ALIGN.CENTER
    arrow = slide_textbox(s10, Inches(5.85), Inches(6.1), Inches(1.8), Inches(0.8), "→ Zukunft", 24, True)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = COLOR_WHITE
    arrow.line.color.rgb = COLOR_TEXT
    arrow.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_transition(s10, "wheel")
    add_notes(
        s10,
        "Aus der Analyse leitet die Facharbeit drei zentrale Reformansätze ab. "
        "Erstens: mehr Lebenskompetenzen im Curriculum, also Inhalte mit direktem Alltagsbezug wie finanzielle "
        "Grundbildung oder Medienkompetenz. "
        "Zweitens: eine spätere Leistungsdifferenzierung, damit Bildungsentscheidungen weniger früh festgelegt "
        "werden und Lernende mehr Entwicklungszeit erhalten. "
        "Drittens: alternative Bewertungsformen, die Lernprozesse stärker berücksichtigen und den Fokus nicht "
        "ausschließlich auf Prüfungsnoten legen. "
        "Diese Ansätze sollen nicht Leistung abschaffen, sondern Schule ausgewogener gestalten. "
        "Zum Abschluss fasse ich nun die Kernbotschaft der Arbeit in drei Aussagen zusammen.",
    )

    # Slide 11
    s11 = prs.slides.add_slide(layout)
    decorate_background(s11, 1)
    add_title(s11, "Schule beeinflusst mehr als nur Noten.", top=Inches(0.8), size=50)
    final = add_round_rect(s11, Inches(1.0), Inches(2.0), Inches(11.2), Inches(3.6), COLOR_WHITE, line_color=COLOR_TEXT)
    final.text_frame.text = (
        "✓ Strukturen wirken auf das Wohlbefinden.\n"
        "✓ Leistungsdruck und Schulklima spielen eine zentrale Rolle.\n"
        "✓ Reformen könnten Bildung und Glück stärker miteinander verbinden."
    )
    for p in final.text_frame.paragraphs:
        p.font.name = "Avenir Next"
        p.font.size = Pt(31)
        p.font.color.rgb = COLOR_TEXT
        p.font.bold = True
    thanks = slide_textbox(s11, Inches(3.85), Inches(6.25), Inches(5.6), Inches(0.9), "Danke für eure Aufmerksamkeit!", 30, True)
    thanks.fill.solid()
    thanks.fill.fore_color.rgb = COLOR_BG_BLUE
    thanks.line.color.rgb = COLOR_TEXT
    thanks.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_transition(s11, "fade")
    add_notes(
        s11,
        "Damit komme ich zum Fazit. "
        "Die Leitfrage kann klar beantwortet werden: Strukturelle Unterschiede im deutschen Schulsystem "
        "beeinflussen das Wohlbefinden von Schülerinnen und Schülern spürbar. "
        "Die Facharbeit zeigt, dass insbesondere frühe Selektion, Leistungsdruck und Schulklima entscheidende "
        "Wirkfaktoren sind. Gleichzeitig wird deutlich, dass Schule auch Chancen bietet, wenn Strukturen "
        "unterstützend gestaltet werden. "
        "Die diskutierten Reformansätze zeigen, wie Bildungsgerechtigkeit, Leistungsanspruch und Wohlbefinden "
        "besser verbunden werden können. "
        "Damit endet meine Präsentation. Vielen Dank fürs Zuhören – ich freue mich auf Fragen und eure Perspektiven.",
    )

    output_dir = Path("output")
    output_dir.mkdir(exist_ok=True)
    out_file = output_dir / "Seminarfach_Glueck_Keynote_Style.pptx"
    prs.save(out_file)
    print(f"Created: {out_file}")


def slide_textbox(slide, left, top, width, height, text: str, size: int, bold: bool):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Avenir Next"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = COLOR_TEXT
    return box


def slide_circle(slide, left, top, diameter, color):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, left, top, diameter, diameter)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.color.rgb = COLOR_TEXT
    shp.line.width = Pt(1.5)
    return shp


if __name__ == "__main__":
    main()
