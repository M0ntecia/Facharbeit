#!/usr/bin/env python3
"""Generate Keynote-style presentation from Facharbeit content."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Pastell palette
HELLBLAU = RGBColor(0xCF, 0xEF, 0xFF)
PASTELLROSA = RGBColor(0xFF, 0xD9, 0xE8)
HELLGELB = RGBColor(0xFF, 0xF7, 0xC7)
WEISS = RGBColor(0xFF, 0xFF, 0xFF)
DUNKELGRAU = RGBColor(0x2D, 0x2D, 0x2D)
MITTELGRAU = RGBColor(0x66, 0x66, 0x66)
HELLGRUEN = RGBColor(0xD4, 0xF5, 0xE0)
AKZENTBLAU = RGBColor(0x5B, 0xA8, 0xD4)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

SPEAKER_NOTES = {
    1: (
        "Guten Tag zusammen! Mein Name ist Svea Timphus, und heute möchte ich euch meine Facharbeit "
        "aus dem Seminarfachkurs Glück vorstellen. Das Thema lautet: Wie Schule wirkt – strukturelle "
        "Unterschiede und ihr Effekt auf das Wohlbefinden von Schülerinnen und Schülern. "
        "Schule ist für uns alle ein zentraler Ort: Wir verbringen dort einen großen Teil unseres "
        "Alltags. Gleichzeitig ist sie aber nicht nur ein Ort des Lernens, sondern auch ein Ort der "
        "Leistungsbewertung und sozialen Vergleiche. Genau diesen Zusammenhang zwischen Schulstruktur "
        "und Wohlbefinden möchte ich in den nächsten fünfzehn Minuten mit euch erkunden."
    ),
    2: (
        "Warum ist dieses Thema überhaupt wichtig? Schule prägt unseren Alltag in einem Alter, in dem "
        "wir uns noch stark entwickeln. Sie vermittelt nicht nur Wissen, sondern beeinflusst auch "
        "Selbstwertgefühl, Motivation und Zukunftsperspektiven. In Deutschland nimmt der Leistungsdruck "
        "zu – Studien des Robert Koch-Instituts und des Deutschen Jugendinstituts zeigen, dass sich viele "
        "Jugendliche durch schulische Anforderungen stark belastet fühlen. Bildung ist zugleich ein "
        "entscheidender Faktor für soziale Teilhabe und berufliche Chancen. Deshalb ist Wohlbefinden "
        "kein Nebenthema, sondern ein wichtiger Bestandteil erfolgreicher Bildung. Meine zentrale "
        "Leitfrage lautet daher: Inwiefern beeinflussen strukturelle Unterschiede im deutschen "
        "Schulsystem das Wohlbefinden von Schülerinnen und Schülern?"
    ),
    3: (
        "Bevor wir ins Detail gehen, klären wir die Grundlagen. Unter Strukturmerkmalen verstehe ich "
        "institutionelle Rahmenbedingungen wie Schulformen, das Notensystem, Versetzungsordnungen und "
        "Abschlussprüfungen. Wohlbefinden umfasst sowohl psychische Gesundheit – also Stress, Angst "
        "oder Schlafprobleme – als auch soziale Beziehungen und Lebenszufriedenheit. Die WHO definiert "
        "Gesundheit nicht nur als Abwesenheit von Krankheit, sondern als vollständiges körperliches, "
        "geistiges und soziales Wohlbefinden. Drittens spielt Bildungsgerechtigkeit eine Rolle: Es geht "
        "darum, ob alle Schülerinnen und Schüler faire Chancen haben – unabhängig von ihrer sozialen "
        "Herkunft. Diese drei Begriffe bilden den Rahmen für meine gesamte Analyse."
    ),
    4: (
        "Das deutsche Schulsystem ist durch seine Mehrgliedrigkeit geprägt. Nach der Grundschule werden "
        "Schülerinnen und Schüler auf verschiedene Schulformen verteilt. Das Gymnasium bereitet auf das "
        "Abitur und ein Studium vor – mit hohen Anforderungen und häufigem Prüfungsdruck. Die Realschule "
        "verbindet allgemeinbildende Inhalte mit stärkerer Praxisorientierung und führt zum Mittleren "
        "Schulabschluss. Die Hauptschule war traditionell stärker berufsorientiert, wird aber häufig "
        "mit geringeren gesellschaftlichen Chancen und Stigmatisierung verbunden. Die Gesamtschule "
        "ermöglicht längeres gemeinsames Lernen mit individueller Förderung. Und berufliche Schulen "
        "sind Teil des dualen Ausbildungssystems, das international als besonderes Merkmal Deutschlands "
        "gilt. Jede Schulform prägt den schulischen Alltag und die Zukunftsperspektiven anders."
    ),
    5: (
        "Ein zentraler Faktor für das Wohlbefinden ist der Leistungsdruck. Links seht ihr die "
        "Auslöser: Noten entscheiden über Versetzungen und Abschlüsse, Prüfungen wie Klassenarbeiten "
        "und Tests erzeugen Stress, die Vergleichskultur fördert Konkurrenzdenken, und hohe "
        "Erwartungen von Eltern, Lehrkräften und der Gesellschaft erhöhen den Druck. Rechts die "
        "möglichen Folgen: Schulischer Stress gehört laut WHO zu den häufigsten Belastungsfaktoren "
        "im Jugendalter. Er kann Schlafprobleme, Kopfschmerzen und Konzentrationsschwierigkeiten "
        "verursachen. Viele Schülerinnen und Schüler entwickeln Unsicherheit oder ein geringeres "
        "Selbstwertgefühl. Wichtig ist: Leistungsdruck kann Motivation fördern, aber auch das "
        "Wohlbefinden deutlich beeinträchtigen – je nachdem, wie wir die Anforderungen erleben."
    ),
    6: (
        "Neben Leistungsdruck spielt das Schulklima eine entscheidende Rolle. Das Lehrer-Schüler-"
        "Verhältnis ist einer der wichtigsten Faktoren: Schülerinnen und Schüler fühlen sich besser, "
        "wenn sie ernst genommen werden, Unterstützung erhalten und ein respektvolles Verhältnis "
        "besteht. Mitbestimmung stärkt das Gefühl von Selbstwirksamkeit – etwa durch Schülervertretung, "
        "Klassenräte oder Schulprojekte. Und eine gute Feedbackkultur geht über reine Noten hinaus: "
        "Qualitative Rückmeldungen und formative Bewertungen können Lernprozesse begleiten und "
        "individuelle Entwicklung stärker berücksichtigen. Der nationale Bildungsbericht betont, "
        "dass ein positives Schulklima mit höherer Lernmotivation und besserer psychischer Gesundheit "
        "verbunden ist."
    ),
    7: (
        "Jetzt möchte ich euch kurz einbeziehen! Stellt euch vor, ich frage euch: Was beeinflusst "
        "euer Wohlbefinden in der Schule am stärksten? Notendruck, Lehrkräfte, Mitschüler oder "
        "Zeitstress? Zeigt mir bitte kurz per Handzeichen, was für euch am wichtigsten ist. "
        "[Pause für Umfrage] Interessant – und genau das zeigt, wie vielfältig die Faktoren sind. "
        "In meiner Facharbeit wird deutlich, dass alle diese Aspekte eine wichtige Rolle spielen: "
        "Leistungsbewertungen und Prüfungen erzeugen Druck, das Verhältnis zu Lehrkräften prägt "
        "das Schulklima, Mitschüler beeinflussen soziale Eingebundenheit, und Zeitstress gehört "
        "zu den häufig genannten Belastungsfaktoren. Wohlbefinden entsteht also aus dem "
        "Zusammenspiel vieler Faktoren – nicht aus einem einzelnen."
    ),
    8: (
        "Um das deutsche System einzuordnen, schauen wir international. In Deutschland erfolgt die "
        "Aufteilung auf Schulformen früh – meist nach der vierten Klasse – und Noten spielen eine "
        "zentrale Rolle. Finnland verfolgt ein anderes Modell: Kinder lernen bis zur neunten Klasse "
        "gemeinsam, die Leistungsdifferenzierung erfolgt deutlich später. Dort werden in frühen "
        "Jahren seltener standardisierte Tests eingesetzt, stattdessen stehen individuelle "
        "Lernfortschritte im Vordergrund. Angelsächsische Systeme wie in Großbritannien oder den USA "
        "bieten mehr Wahlmöglichkeiten im Unterricht und legen stärkeren Fokus auf projektorientiertes "
        "und praxisorientiertes Lernen. Aber: Kein System ist perfekt. Jedes muss unterschiedliche "
        "gesellschaftliche Anforderungen erfüllen – von akademischer Vorbereitung bis zur beruflichen "
        "Ausbildung."
    ),
    9: (
        "Eine kritische Bewertung zeigt Stärken und Schwächen des deutschen Systems. Zu den Stärken "
        "gehören die verschiedenen Bildungswege, die eine gezielte Förderung ermöglichen, das duale "
        "Ausbildungssystem, das den Übergang in den Arbeitsmarkt erleichtert, und hohe akademische "
        "Standards im Gymnasium. Auf der anderen Seite stehen Schwächen: die frühe Selektion, die "
        "Bildungsentscheidungen schon in jungen Jahren festlegt, der hohe Leistungsdruck durch Noten "
        "und Prüfungen, und soziale Ungleichheiten – denn der Bildungserfolg in Deutschland hängt "
        "laut OECD vergleichsweise stark vom sozioökonomischen Hintergrund ab. Diese Spannung "
        "zwischen Differenzierung und Gerechtigkeit prägt die gesamte bildungspolitische Debatte."
    ),
    10: (
        "Was könnte sich ändern? In der Bildungsforschung werden drei Reformansätze diskutiert. "
        "Erstens: Mehr Lebenskompetenzen im Curriculum – etwa finanzielle Grundbildung, "
        "Medienkompetenz und praktische Alltagsfähigkeiten. Zweitens: Spätere Leistungsdifferenzierung, "
        "wie sie in skandinavischen Ländern praktiziert wird, um Schülerinnen und Schülern mehr Zeit "
        "für ihre Entwicklung zu geben. Drittens: Alternative Bewertungsformen wie Portfolioarbeit, "
        "ausführliches Lernfeedback oder projektbasierte Bewertung – statt ausschließlich Noten und "
        "Prüfungsergebnisse. Diese Ansätze sollen Bildungsqualität und Schülerwohlbefinden stärker "
        "miteinander verbinden."
    ),
    11: (
        "Zum Fazit: Schule beeinflusst mehr als nur Noten. Strukturelle Merkmale wie Schulformen, "
        "Bewertungssysteme und Selektionsmechanismen wirken direkt auf das Wohlbefinden. Leistungsdruck "
        "und Schulklima spielen dabei eine zentrale Rolle – einerseits durch Noten und Prüfungen, "
        "andererseits durch Beziehungen und Partizipation. Und Reformen könnten Bildung und Glück "
        "stärker verbinden, etwa durch spätere Differenzierung, Lebenskompetenzen und neue "
        "Bewertungsformen. Wohlbefinden hängt nicht nur von uns als Individuen ab, sondern in hohem "
        "Maße von den institutionellen Rahmenbedingungen. Vielen Dank für eure Aufmerksamkeit – ich "
        "freue mich auf eure Fragen!"
    ),
}

TRANSITIONS = [
    "fade",      # Folie 1
    "zoom",      # Folie 2
    "cover",     # Folie 3 – Move In
    "fade",      # Folie 4 – Magic Move (in Keynote manuell)
    "zoom",      # Folie 5 – Scale
    "push",      # Folie 6
    "dissolve",  # Folie 7
    "zoom",      # Folie 8 – Pop (in Keynote manuell)
    "push",      # Folie 9 – Slide
    "fade",      # Folie 10 – Rotate (in Keynote manuell)
    "fade",      # Folie 11
]


def hex_rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def add_rounded_rect(slide, left, top, width, height, fill, line=None, radius_hint=True):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if radius_hint and hasattr(shape, "adjustments") and len(shape.adjustments) > 0:
        shape.adjustments[0] = 0.12
    return shape


def add_shadow_effect(shape):
    try:
        shape.shadow.inherit = False
        shape.shadow.visible = True
    except Exception:
        pass


def set_text(
    shape,
    text: str,
    size: int = 18,
    bold: bool = False,
    color: RGBColor = DUNKELGRAU,
    align=PP_ALIGN.LEFT,
):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Helvetica Neue"
    p.alignment = align
    return tf


def add_bullets(shape, items: list[str], size: int = 16, color: RGBColor = DUNKELGRAU):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = "Helvetica Neue"
        p.space_after = Pt(8)
        p.bullet = True


def add_title(slide, text: str, top=Inches(0.45), size=40):
    box = slide.shapes.add_textbox(Inches(0.7), top, Inches(11.8), Inches(0.9))
    set_text(box, text, size=size, bold=True, color=DUNKELGRAU)
    return box


def add_notes(slide, note: str):
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = note


def create_illustration_school(path: Path):
    img = Image.new("RGBA", (800, 500), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle([250, 120, 550, 380], radius=20, fill=(207, 239, 255, 255))
    draw.polygon([(300, 120), (500, 40), (700, 120)], fill=(255, 217, 232, 255))
    draw.rectangle([360, 200, 440, 380], fill=(255, 247, 199, 255))
    draw.rectangle([460, 200, 540, 380], fill=(255, 247, 199, 255))
    for x in (380, 480):
        draw.ellipse([x, 150, x + 40, x + 40], fill=(255, 217, 232, 255))
    draw.ellipse([120, 300, 180, 420], fill=(212, 245, 224, 255))
    draw.ellipse([620, 300, 680, 420], fill=(255, 217, 232, 255))
    img.save(path, "PNG")


def create_illustration_scale(path: Path):
    img = Image.new("RGBA", (600, 400), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle([250, 180, 350, 360], fill=(207, 239, 255, 255))
    draw.line([(100, 180), (500, 180)], fill=(45, 45, 45, 255), width=6)
    draw.ellipse([90, 160, 130, 200], fill=(255, 217, 232, 255))
    draw.ellipse([470, 160, 510, 200], fill=(207, 239, 255, 255))
    draw.text((105, 120), "Bildung", fill=(45, 45, 45, 255))
    draw.text((455, 120), "Wohlbefinden", fill=(45, 45, 45, 255))
    img.save(path, "PNG")


def create_illustration_mindmap(path: Path):
    img = Image.new("RGBA", (700, 500), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    draw.ellipse([280, 200, 420, 300], fill=(255, 247, 199, 255))
    draw.text((310, 235), "Schule", fill=(45, 45, 45, 255))
    for pos, label, color in [
        ((80, 80), "Struktur", (207, 239, 255, 255)),
        ((500, 80), "Wohlbefinden", (255, 217, 232, 255)),
        ((290, 400), "Gerechtigkeit", (212, 245, 224, 255)),
    ]:
        draw.rounded_rectangle([pos[0], pos[1], pos[0] + 160, pos[1] + 60], radius=15, fill=color)
        draw.text((pos[0] + 20, pos[1] + 18), label, fill=(45, 45, 45, 255))
    img.save(path, "PNG")


def create_illustration_stress(path: Path):
    img = Image.new("RGBA", (400, 500), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    draw.ellipse([120, 80, 280, 240], fill=(255, 217, 232, 255))
    draw.rounded_rectangle([140, 240, 260, 420], radius=20, fill=(207, 239, 255, 255))
    draw.line([(100, 60), (130, 90)], fill=(45, 45, 45, 255), width=4)
    draw.line([(300, 60), (270, 90)], fill=(45, 45, 45, 255), width=4)
    img.save(path, "PNG")


def create_illustration_conversation(path: Path):
    img = Image.new("RGBA", (600, 400), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    draw.ellipse([80, 100, 180, 200], fill=(207, 239, 255, 255))
    draw.rounded_rectangle([60, 200, 200, 360], radius=15, fill=(255, 247, 199, 255))
    draw.ellipse([380, 80, 500, 200], fill=(255, 217, 232, 255))
    draw.rounded_rectangle([360, 200, 520, 360], radius=15, fill=(212, 245, 224, 255))
    draw.rounded_rectangle([220, 140, 320, 200], radius=20, fill=(255, 255, 255, 220), outline=(91, 168, 212, 255))
    img.save(path, "PNG")


def create_illustration_future(path: Path):
    img = Image.new("RGBA", (700, 300), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    draw.line([(80, 220), (620, 80)], fill=(91, 168, 212, 255), width=8)
    draw.polygon([(620, 80), (590, 95), (605, 110)], fill=(91, 168, 212, 255))
    for x, c in [(150, (207, 239, 255, 255)), (350, (255, 217, 232, 255)), (550, (255, 247, 199, 255))]:
        draw.ellipse([x - 50, 140, x + 50, 240], fill=c)
    img.save(path, "PNG")


def create_illustration_calm(path: Path):
    img = Image.new("RGBA", (800, 300), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle([100, 80, 700, 240], radius=40, fill=(207, 239, 255, 180))
    draw.ellipse([200, 120, 280, 200], fill=(255, 217, 232, 200))
    draw.ellipse([520, 120, 600, 200], fill=(255, 247, 199, 200))
    img.save(path, "PNG")


def slide_background(slide, color: RGBColor = WEISS):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def build_slide_1(prs, assets: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_background(slide)
    add_rounded_rect(slide, Inches(0), Inches(0), Inches(5.5), SLIDE_H, HELLBLAU)
    add_rounded_rect(slide, Inches(8.8), Inches(0), Inches(4.5), Inches(2.2), PASTELLROSA)
    add_rounded_rect(slide, Inches(9.5), Inches(5.2), Inches(3.5), Inches(2), HELLGELB)

    title = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(5.8), Inches(2.2))
    tf = title.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "Wie Schule wirkt:"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = DUNKELGRAU
    p1.font.name = "Helvetica Neue"
    p2 = tf.add_paragraph()
    p2.text = "Strukturelle Unterschiede und ihr Effekt auf das Wohlbefinden von Schülerinnen und Schülern"
    p2.font.size = Pt(22)
    p2.font.color.rgb = MITTELGRAU
    p2.font.name = "Helvetica Neue"
    p2.space_before = Pt(12)

    sub = slide.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(5), Inches(0.6))
    set_text(sub, "Seminarfachkurs Glück", size=20, color=AKZENTBLAU)

    img_path = assets / "school.png"
    create_illustration_school(img_path)
    slide.shapes.add_picture(str(img_path), Inches(6.2), Inches(1.2), width=Inches(6.5))
    add_notes(slide, SPEAKER_NOTES[1])


def build_slide_2(prs, assets: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_background(slide)
    add_title(slide, "Warum ist das Thema wichtig?")

    img_path = assets / "scale.png"
    create_illustration_scale(img_path)
    slide.shapes.add_picture(str(img_path), Inches(8.5), Inches(1.2), width=Inches(4))

    card = add_rounded_rect(slide, Inches(0.7), Inches(1.3), Inches(7.5), Inches(3.2), WEISS, AKZENTBLAU)
    add_shadow_effect(card)
    bullets = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(7), Inches(2.6))
    add_bullets(
        bullets,
        [
            "Schule prägt Alltag und ist ein prägender Lebensraum",
            "Leistungsdruck und psychische Belastungen nehmen zu",
            "Bildung beeinflusst soziale Teilhabe und Zukunft",
            "Wohlbefinden ist wichtiger Bestandteil erfolgreicher Bildung",
        ],
        size=17,
    )

    box = add_rounded_rect(slide, Inches(0.7), Inches(4.8), Inches(11.8), Inches(1.8), HELLGELB)
    add_shadow_effect(box)
    q = slide.shapes.add_textbox(Inches(1.0), Inches(5.1), Inches(11.2), Inches(1.3))
    set_text(
        q,
        '"Inwiefern beeinflussen strukturelle Unterschiede im deutschen Schulsystem das Wohlbefinden von Schülerinnen und Schülern?"',
        size=17,
        bold=True,
        color=DUNKELGRAU,
        align=PP_ALIGN.CENTER,
    )
    add_notes(slide, SPEAKER_NOTES[2])


def build_slide_3(prs, assets: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_background(slide, HELLBLAU)
    add_title(slide, "Grundlagen")

    img_path = assets / "mindmap.png"
    create_illustration_mindmap(img_path)
    slide.shapes.add_picture(str(img_path), Inches(0.5), Inches(1.1), width=Inches(4.5))

    columns = [
        ("Strukturmerkmale", HELLBLAU, ["Schulformen", "Notensystem", "Versetzungen", "Abschlüsse"]),
        ("Wohlbefinden", PASTELLROSA, ["Psychische Gesundheit", "Soziale Beziehungen", "Lebenszufriedenheit"]),
        (
            "Bildungsgerechtigkeit",
            HELLGELB,
            ["Faire Chancen unabhängig von sozialer Herkunft"],
        ),
    ]
    x = Inches(5.2)
    for title, color, items in columns:
        card = add_rounded_rect(slide, x, Inches(1.3), Inches(2.5), Inches(5.2), WEISS, AKZENTBLAU)
        add_shadow_effect(card)
        hdr = slide.shapes.add_textbox(x + Inches(0.2), Inches(1.5), Inches(2.1), Inches(0.5))
        set_text(hdr, title, size=16, bold=True, color=DUNKELGRAU, align=PP_ALIGN.CENTER)
        accent = add_rounded_rect(slide, x + Inches(0.3), Inches(2.0), Inches(1.9), Inches(0.08), color)
        body = slide.shapes.add_textbox(x + Inches(0.2), Inches(2.2), Inches(2.1), Inches(3.8))
        add_bullets(body, items, size=14)
        x += Inches(2.7)
    add_notes(slide, SPEAKER_NOTES[3])


def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_background(slide)
    add_title(slide, "Das deutsche Schulsystem")

    schools = [
        ("Gymnasium", "Abitur, hohe Anforderungen, Prüfungsdruck", "🎓", HELLBLAU),
        ("Realschule", "Mittlerer Abschluss, Praxisorientierung", "📐", PASTELLROSA),
        ("Hauptschule", "Berufsorientierung, Stigmatisierung möglich", "🔧", HELLGELB),
        ("Gesamtschule", "Länger gemeinsam lernen, individuelle Förderung", "🤝", HELLGRUEN),
        ("Berufliche Schule", "Duales Ausbildungssystem, Praxis + Theorie", "🏭", HELLBLAU),
    ]
    x = Inches(0.5)
    for name, desc, icon, color in schools:
        card = add_rounded_rect(slide, x, Inches(1.5), Inches(2.35), Inches(5.0), color)
        add_shadow_effect(card)
        ic = slide.shapes.add_textbox(x + Inches(0.15), Inches(1.7), Inches(2), Inches(0.6))
        set_text(ic, icon, size=28, align=PP_ALIGN.CENTER)
        nm = slide.shapes.add_textbox(x + Inches(0.15), Inches(2.3), Inches(2.05), Inches(0.5))
        set_text(nm, name, size=15, bold=True, align=PP_ALIGN.CENTER)
        ds = slide.shapes.add_textbox(x + Inches(0.15), Inches(2.9), Inches(2.05), Inches(3.2))
        tf = ds.text_frame
        tf.word_wrap = True
        set_text(ds, desc, size=13, color=MITTELGRAU, align=PP_ALIGN.CENTER)
        x += Inches(2.5)
    add_notes(slide, SPEAKER_NOTES[4])


def build_slide_5(prs, assets: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_background(slide)
    add_title(slide, "Leistungsdruck und Wohlbefinden")

    img_path = assets / "stress.png"
    create_illustration_stress(img_path)
    slide.shapes.add_picture(str(img_path), Inches(10.8), Inches(1.2), width=Inches(2))

    left = add_rounded_rect(slide, Inches(0.7), Inches(1.3), Inches(5.5), Inches(3.5), PASTELLROSA)
    add_shadow_effect(left)
    lt = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5), Inches(0.5))
    set_text(lt, "Leistungsdruck", size=20, bold=True)
    lb = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.8), Inches(2.5))
    add_bullets(lb, ["Noten", "Prüfungen", "Konkurrenz", "Erwartungen"], size=16)

    right = add_rounded_rect(slide, Inches(6.5), Inches(1.3), Inches(5.5), Inches(3.5), HELLBLAU)
    add_shadow_effect(right)
    rt = slide.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(5), Inches(0.5))
    set_text(rt, "Folgen", size=20, bold=True)
    rb = slide.shapes.add_textbox(Inches(6.8), Inches(2.0), Inches(4.8), Inches(2.5))
    add_bullets(rb, ["Stress", "Schlafprobleme", "Unsicherheit", "Geringeres Selbstwertgefühl"], size=16)

    quote = add_rounded_rect(slide, Inches(0.7), Inches(5.2), Inches(11.8), Inches(1.5), HELLGELB)
    add_shadow_effect(quote)
    qt = slide.shapes.add_textbox(Inches(1.0), Inches(5.5), Inches(11.2), Inches(1))
    set_text(
        qt,
        '"Leistungsdruck kann Motivation fördern, aber auch das Wohlbefinden beeinträchtigen."',
        size=17,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_notes(slide, SPEAKER_NOTES[5])


def build_slide_6(prs, assets: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_background(slide)
    add_title(slide, "Schulklima")

    img_path = assets / "conversation.png"
    create_illustration_conversation(img_path)
    slide.shapes.add_picture(str(img_path), Inches(9.8), Inches(1.0), width=Inches(3))

    cards = [
        (
            "Lehrer-Schüler-Verhältnis",
            HELLBLAU,
            ["Ernst genommen werden", "Unterstützung erhalten", "Respektvolles Miteinander"],
        ),
        (
            "Mitbestimmung",
            PASTELLROSA,
            ["Schülervertretung & Klassenrat", "Selbstwirksamkeit stärken", "Aktive Beteiligung"],
        ),
        (
            "Feedbackkultur",
            HELLGELB,
            ["Qualitative Rückmeldungen", "Formative Bewertung", "Individuelle Entwicklung"],
        ),
    ]
    x = Inches(0.6)
    for title, color, items in cards:
        card = add_rounded_rect(slide, x, Inches(1.4), Inches(3.8), Inches(5.0), color)
        add_shadow_effect(card)
        hdr = slide.shapes.add_textbox(x + Inches(0.2), Inches(1.6), Inches(3.4), Inches(0.7))
        set_text(hdr, title, size=17, bold=True, align=PP_ALIGN.CENTER)
        body = slide.shapes.add_textbox(x + Inches(0.25), Inches(2.5), Inches(3.3), Inches(3.5))
        add_bullets(body, items, size=14)
        x += Inches(4.1)
    add_notes(slide, SPEAKER_NOTES[6])


def build_slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_background(slide, PASTELLROSA)
    add_title(slide, "Kurze Umfrage")

    qbox = add_rounded_rect(slide, Inches(1.5), Inches(1.5), Inches(10.3), Inches(1.5), WEISS, AKZENTBLAU)
    add_shadow_effect(qbox)
    q = slide.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(9.7), Inches(1))
    set_text(
        q,
        '"Was beeinflusst euer Wohlbefinden in der Schule am stärksten?"',
        size=22,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    options = [
        ("📚 Notendruck", HELLBLAU),
        ("👩‍🏫 Lehrkräfte", PASTELLROSA),
        ("👥 Mitschüler", HELLGELB),
        ("🕒 Zeitstress", HELLGRUEN),
    ]
    x = Inches(1.0)
    for label, color in options:
        btn = add_rounded_rect(slide, x, Inches(3.5), Inches(2.7), Inches(1.2), color)
        add_shadow_effect(btn)
        t = slide.shapes.add_textbox(x + Inches(0.15), Inches(3.8), Inches(2.4), Inches(0.7))
        set_text(t, label, size=16, bold=True, align=PP_ALIGN.CENTER)
        x += Inches(2.95)

    hint = slide.shapes.add_textbox(Inches(1.5), Inches(5.0), Inches(10), Inches(0.5))
    set_text(hint, "Kurzes Handzeichen", size=18, color=MITTELGRAU, align=PP_ALIGN.CENTER)

    reveal = add_rounded_rect(slide, Inches(1.5), Inches(5.7), Inches(10.3), Inches(1.0), WEISS, AKZENTBLAU)
    add_shadow_effect(reveal)
    rt = slide.shapes.add_textbox(Inches(1.8), Inches(5.95), Inches(9.7), Inches(0.6))
    set_text(
        rt,
        "Alle Faktoren spielen laut Facharbeit eine wichtige Rolle.",
        size=17,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_notes(slide, SPEAKER_NOTES[7])


def build_slide_8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_background(slide)
    add_title(slide, "Internationaler Vergleich")

    countries = [
        (
            "Deutschland",
            HELLBLAU,
            ["Frühe Aufteilung nach der Grundschule", "Viele Noten und Prüfungen", "Starke Leistungsselektion"],
        ),
        (
            "Finnland",
            PASTELLROSA,
            ["Längeres gemeinsames Lernen", "Spätere Leistungsdifferenzierung", "Weniger standardisierte Tests"],
        ),
        (
            "Angelsächsische Systeme",
            HELLGELB,
            ["Mehr Praxisorientierung", "Mehr Wahlmöglichkeiten", "Projektorientiertes Lernen"],
        ),
    ]
    x = Inches(0.7)
    for name, color, items in countries:
        card = add_rounded_rect(slide, x, Inches(1.4), Inches(3.8), Inches(4.5), color)
        add_shadow_effect(card)
        hdr = slide.shapes.add_textbox(x + Inches(0.2), Inches(1.6), Inches(3.4), Inches(0.6))
        set_text(hdr, name, size=18, bold=True, align=PP_ALIGN.CENTER)
        body = slide.shapes.add_textbox(x + Inches(0.25), Inches(2.4), Inches(3.3), Inches(3.2))
        add_bullets(body, items, size=14)
        x += Inches(4.1)

    foot = add_rounded_rect(slide, Inches(2.5), Inches(6.2), Inches(8.3), Inches(0.7), WEISS, AKZENTBLAU)
    ft = slide.shapes.add_textbox(Inches(2.7), Inches(6.35), Inches(7.9), Inches(0.5))
    set_text(ft, "Kein System ist perfekt.", size=18, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, SPEAKER_NOTES[8])


def build_slide_9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_background(slide)
    add_title(slide, "Stärken und Schwächen")

    left = add_rounded_rect(slide, Inches(0.7), Inches(1.4), Inches(5.8), Inches(5.2), HELLGRUEN)
    add_shadow_effect(left)
    lt = slide.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.2), Inches(0.6))
    set_text(lt, "Stärken", size=24, bold=True)
    lb = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(5.2), Inches(3.5))
    add_bullets(
        lb,
        [
            "Verschiedene Bildungswege ermöglichen gezielte Förderung",
            "Duales Ausbildungssystem erleichtert den Berufseinstieg",
            "Hohe akademische Standards im Gymnasium",
        ],
        size=16,
    )

    right = add_rounded_rect(slide, Inches(6.8), Inches(1.4), Inches(5.8), Inches(5.2), PASTELLROSA)
    add_shadow_effect(right)
    rt = slide.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.2), Inches(0.6))
    set_text(rt, "Schwächen", size=24, bold=True)
    rb = slide.shapes.add_textbox(Inches(7.1), Inches(2.5), Inches(5.2), Inches(3.5))
    add_bullets(
        rb,
        [
            "Frühe Selektion nach der Grundschule",
            "Hoher Leistungsdruck durch Noten und Prüfungen",
            "Soziale Ungleichheiten und Segregation",
        ],
        size=16,
    )
    add_notes(slide, SPEAKER_NOTES[9])


def build_slide_10(prs, assets: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_background(slide)
    add_title(slide, "Reformansätze")

    img_path = assets / "future.png"
    create_illustration_future(img_path)
    slide.shapes.add_picture(str(img_path), Inches(0.5), Inches(5.5), width=Inches(12.3))

    reforms = [
        (
            "Mehr Lebenskompetenzen",
            HELLBLAU,
            "Finanzbildung, Medienkompetenz, praktische Alltagsfähigkeiten",
        ),
        (
            "Spätere Leistungsdifferenzierung",
            PASTELLROSA,
            "Mehr Zeit für Entwicklung, wie in skandinavischen Modellen",
        ),
        (
            "Alternative Bewertungsformen",
            HELLGELB,
            "Portfolioarbeit, Lernfeedback, projektbasierte Bewertung",
        ),
    ]
    x = Inches(0.8)
    for title, color, desc in reforms:
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x, Inches(1.8), Inches(3.5), Inches(2.2))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        add_shadow_effect(circle)
        ct = slide.shapes.add_textbox(x + Inches(0.15), Inches(2.2), Inches(3.2), Inches(1.2))
        tf = ct.text_frame
        tf.word_wrap = True
        set_text(ct, title, size=15, bold=True, align=PP_ALIGN.CENTER)
        ds = slide.shapes.add_textbox(x + Inches(0.1), Inches(4.2), Inches(3.4), Inches(1.2))
        tf2 = ds.text_frame
        tf2.word_wrap = True
        set_text(ds, desc, size=13, color=MITTELGRAU, align=PP_ALIGN.CENTER)
        x += Inches(4.0)
    add_notes(slide, SPEAKER_NOTES[10])


def build_slide_11(prs, assets: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_background(slide, HELLBLAU)
    add_rounded_rect(slide, Inches(0), Inches(5.5), SLIDE_W, Inches(2), PASTELLROSA)

    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(1.2))
    set_text(title, "Schule beeinflusst mehr als nur Noten.", size=36, bold=True, align=PP_ALIGN.CENTER)

    points = [
        "Strukturen wirken auf das Wohlbefinden.",
        "Leistungsdruck und Schulklima spielen eine zentrale Rolle.",
        "Reformen könnten Bildung und Glück stärker miteinander verbinden.",
    ]
    y = Inches(2.2)
    for point in points:
        row = add_rounded_rect(slide, Inches(1.5), y, Inches(10.3), Inches(0.75), WEISS, AKZENTBLAU)
        add_shadow_effect(row)
        t = slide.shapes.add_textbox(Inches(1.8), y + Inches(0.12), Inches(9.7), Inches(0.55))
        set_text(t, f"✓  {point}", size=17, color=DUNKELGRAU)
        y += Inches(0.95)

    img_path = assets / "calm.png"
    create_illustration_calm(img_path)
    slide.shapes.add_picture(str(img_path), Inches(1.5), Inches(5.7), width=Inches(5))

    thanks = slide.shapes.add_textbox(Inches(6.5), Inches(6.2), Inches(6), Inches(0.8))
    set_text(thanks, "Danke für eure Aufmerksamkeit!", size=24, bold=True, color=DUNKELGRAU)
    add_notes(slide, SPEAKER_NOTES[11])


def add_transitions(pptx_path: Path):
    """Inject slide transition XML into PPTX."""
    ns = {
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }
    ET.register_namespace("p", ns["p"])
    ET.register_namespace("a", ns["a"])

    transition_map = {
        "fade": "fade",
        "zoom": "zoom",
        "push": "push",
        "dissolve": "randomBar",
        "cover": "cover",
    }

    buf = io.BytesIO()
    with zipfile.ZipFile(pptx_path, "r") as zin:
        slide_files = sorted(
            [
                n
                for n in zin.namelist()
                if n.startswith("ppt/slides/slide") and n.endswith(".xml") and n.count("/") == 2
            ],
            key=lambda s: int(s.replace("ppt/slides/slide", "").replace(".xml", "")),
        )
        slide_xml = {}
        for i, name in enumerate(slide_files):
            root = ET.fromstring(zin.read(name))
            trans_type = transition_map.get(TRANSITIONS[i], "fade")
            trans = ET.SubElement(root, f"{{{ns['p']}}}transition")
            trans.set("spd", "med")
            ET.SubElement(trans, f"{{{ns['p']}}}{trans_type}")
            slide_xml[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = slide_xml[item.filename] if item.filename in slide_xml else zin.read(item.filename)
                zout.writestr(item, data)

    pptx_path.write_bytes(buf.getvalue())


def main():
    out_dir = Path(__file__).resolve().parent
    assets = out_dir / "assets"
    assets.mkdir(exist_ok=True)
    output = out_dir / "Wie_Schule_wirkt_Keynote.pptx"

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    build_slide_1(prs, assets)
    build_slide_2(prs, assets)
    build_slide_3(prs, assets)
    build_slide_4(prs)
    build_slide_5(prs, assets)
    build_slide_6(prs, assets)
    build_slide_7(prs)
    build_slide_8(prs)
    build_slide_9(prs)
    build_slide_10(prs, assets)
    build_slide_11(prs, assets)

    prs.save(output)
    add_transitions(output)
    print(f"Created: {output}")


if __name__ == "__main__":
    main()
