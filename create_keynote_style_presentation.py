from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


COLORS = {
    "blue": RGBColor(0xCF, 0xEF, 0xFF),
    "pink": RGBColor(0xFF, 0xD9, 0xE8),
    "yellow": RGBColor(0xFF, 0xF7, 0xC7),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
    "dark": RGBColor(0x33, 0x33, 0x33),
    "muted": RGBColor(0x66, 0x66, 0x66),
    "green_soft": RGBColor(0xD8, 0xF3, 0xDC),
}


def add_bg(slide, variant=0):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = COLORS["white"]

    top = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.2), Inches(0.2), Inches(2.2), Inches(0.7)
    )
    top.fill.solid()
    top.fill.fore_color.rgb = COLORS["blue" if variant % 3 == 0 else "pink"]
    top.fill.fore_color.brightness = 0.2
    top.line.fill.background()

    right = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(11.2), Inches(6.6), Inches(2.0), Inches(0.7)
    )
    right.fill.solid()
    right.fill.fore_color.rgb = COLORS["yellow" if variant % 2 == 0 else "blue"]
    right.fill.fore_color.brightness = 0.15
    right.line.fill.background()


def text_box(slide, text, x, y, w, h, size=20, bold=False, color="dark", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Avenir Next"
    run.font.color.rgb = COLORS[color]
    return box


def bullet_box(slide, title, bullets, x, y, w, h, fill_color=None):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS[fill_color]
        shape.fill.fore_color.brightness = 0.08
    else:
        shape.fill.background()
    shape.line.color.rgb = COLORS["muted"]
    shape.line.width = Pt(1.2)

    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.12)
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.runs[0].font.size = Pt(22)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = COLORS["dark"]
    p.runs[0].font.name = "Avenir Next"

    for item in bullets[:5]:
        bp = tf.add_paragraph()
        bp.text = f"• {item}"
        bp.level = 0
        bp.runs[0].font.size = Pt(16)
        bp.runs[0].font.color.rgb = COLORS["dark"]
        bp.runs[0].font.name = "Avenir Next"
        bp.space_after = Pt(2)

    return shape


def set_notes(slide, notes_text):
    slide.notes_slide.notes_text_frame.text = notes_text


def set_transition(slide, effect, attrs=None):
    attrs = attrs or {}
    sld = slide._element
    for element in sld.findall(qn("p:transition")):
        sld.remove(element)

    trans = OxmlElement("p:transition")
    effect_el = OxmlElement(f"p:{effect}")
    for key, value in attrs.items():
        effect_el.set(key, value)
    trans.append(effect_el)

    children = list(sld)
    insert_idx = 1
    if len(children) > 1 and children[1].tag == qn("p:clrMapOvr"):
        insert_idx = 2
    sld.insert(insert_idx, trans)


def card_with_icon(slide, x, y, w, h, emoji, title, lines, fill):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = COLORS[fill]
    card.line.fill.background()

    icon = card.text_frame
    icon.clear()
    p0 = icon.paragraphs[0]
    p0.text = f"{emoji}  {title}"
    p0.runs[0].font.size = Pt(20)
    p0.runs[0].font.bold = True
    p0.runs[0].font.name = "Avenir Next"
    p0.runs[0].font.color.rgb = COLORS["dark"]

    for line in lines[:3]:
        p = icon.add_paragraph()
        p.text = f"• {line}"
        p.runs[0].font.size = Pt(14)
        p.runs[0].font.name = "Avenir Next"
        p.runs[0].font.color.rgb = COLORS["dark"]


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1
    s1 = prs.slides.add_slide(blank)
    add_bg(s1, 1)
    text_box(
        s1,
        "Wie Schule wirkt:",
        0.8,
        0.6,
        8.2,
        0.8,
        size=44,
        bold=True,
    )
    text_box(
        s1,
        "Strukturelle Unterschiede und ihr Effekt\nauf das Wohlbefinden von Schülerinnen und Schülern",
        0.8,
        1.4,
        7.3,
        1.8,
        size=25,
        bold=False,
    )
    text_box(s1, "Seminarfachkurs Glück", 0.8, 3.1, 4.0, 0.5, size=20, color="muted")

    scene = s1.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(0.9), Inches(5.2), Inches(5.8))
    scene.fill.solid()
    scene.fill.fore_color.rgb = COLORS["blue"]
    scene.fill.fore_color.brightness = 0.28
    scene.line.fill.background()
    text_box(s1, "🏫", 9.55, 1.55, 1.4, 1.0, size=64, align=PP_ALIGN.CENTER)
    text_box(s1, "🙂🙂", 8.25, 3.0, 1.8, 0.6, size=36, align=PP_ALIGN.CENTER)
    text_box(s1, "😓😣", 10.5, 3.05, 1.8, 0.6, size=36, align=PP_ALIGN.CENTER)
    text_box(s1, "Lernfreude", 8.35, 3.7, 1.6, 0.4, size=14, color="muted", align=PP_ALIGN.CENTER)
    text_box(s1, "Leistungsdruck", 10.45, 3.7, 1.9, 0.4, size=14, color="muted", align=PP_ALIGN.CENTER)
    set_transition(s1, "fade")

    set_notes(
        s1,
        "Willkommen zu meiner Präsentation im Seminarfachkurs Glück. "
        "Ich untersuche heute, wie Schule nicht nur Leistungen, sondern auch das Wohlbefinden von Schülerinnen und Schülern beeinflusst. "
        "Im Mittelpunkt steht die Frage, welche strukturellen Merkmale des deutschen Schulsystems belastend oder unterstützend wirken. "
        "Die Facharbeit zeigt, dass sich schulische Organisation, Leistungsbewertung und soziale Bedingungen direkt auf Motivation, Selbstwertgefühl und psychische Gesundheit auswirken können. "
        "Diese Gegenüberstellung von zufriedenen und gestressten Lernenden macht das Spannungsfeld sichtbar: Schule kann stärken, aber auch belasten. "
        "Im nächsten Schritt klären wir, warum dieses Thema gesellschaftlich so relevant ist und welche Leitfrage die Arbeit verfolgt.",
    )

    # Slide 2
    s2 = prs.slides.add_slide(blank)
    add_bg(s2, 2)
    text_box(s2, "Warum ist das Thema wichtig?", 0.8, 0.55, 6.5, 0.8, size=39, bold=True)
    bullet_box(
        s2,
        "Relevanz",
        [
            "Schule prägt den Alltag von Kindern und Jugendlichen.",
            "Leistungsdruck und psychische Belastungen werden häufiger diskutiert.",
            "Bildung beeinflusst berufliche Chancen und soziale Teilhabe.",
            "Wohlbefinden ist Teil erfolgreicher Bildung.",
        ],
        0.8,
        1.5,
        6.6,
        3.6,
        fill_color="yellow",
    )
    balance = s2.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.4), Inches(4.6), Inches(3.7))
    balance.fill.solid()
    balance.fill.fore_color.rgb = COLORS["pink"]
    balance.fill.fore_color.brightness = 0.18
    balance.line.fill.background()
    text_box(s2, "⚖️", 9.6, 2.0, 1.5, 1.0, size=66, align=PP_ALIGN.CENTER)
    text_box(s2, "Bildung", 8.55, 3.3, 1.4, 0.4, size=16, bold=True, align=PP_ALIGN.CENTER)
    text_box(s2, "Wohlbefinden", 10.3, 3.3, 1.9, 0.4, size=16, bold=True, align=PP_ALIGN.CENTER)

    q = s2.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(5.35), Inches(11.9), Inches(1.25))
    q.fill.solid()
    q.fill.fore_color.rgb = COLORS["blue"]
    q.fill.fore_color.brightness = 0.08
    q.line.color.rgb = COLORS["muted"]
    tf = q.text_frame
    tf.text = (
        "Inwiefern beeinflussen strukturelle Unterschiede im deutschen Schulsystem "
        "das Wohlbefinden von Schülerinnen und Schülern?"
    )
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].runs[0].font.size = Pt(21)
    tf.paragraphs[0].runs[0].font.bold = True
    tf.paragraphs[0].runs[0].font.name = "Avenir Next"
    tf.paragraphs[0].runs[0].font.color.rgb = COLORS["dark"]
    set_transition(s2, "zoom")

    set_notes(
        s2,
        "Die Facharbeit betont, dass Schule ein zentraler Lebensraum ist: "
        "Ein großer Teil des Alltags findet dort statt, deshalb wirken schulische Strukturen langfristig auf junge Menschen. "
        "Gleichzeitig zeigen Studien, dass schulische Anforderungen häufig als belastend erlebt werden. "
        "Das ist gesellschaftlich relevant, weil Bildung über spätere Chancen mitentscheidet und weil Wohlbefinden eine Voraussetzung für nachhaltiges Lernen ist. "
        "Wenn Schülerinnen und Schüler dauerhaft unter Druck stehen, leidet nicht nur die Gesundheit, sondern oft auch die Lernmotivation. "
        "Daraus entsteht die Leitfrage dieser Arbeit, die wir als roten Faden mitnehmen. "
        "Als Nächstes schauen wir auf die theoretischen Grundlagen, damit die wichtigsten Begriffe klar sind.",
    )

    # Slide 3
    s3 = prs.slides.add_slide(blank)
    add_bg(s3, 3)
    text_box(s3, "Grundlagen", 0.8, 0.55, 4.5, 0.8, size=40, bold=True)
    text_box(s3, "Mindmap zentraler Begriffe", 0.8, 1.2, 4.5, 0.5, size=17, color="muted")

    center = s3.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(5.4), Inches(2.6), Inches(2.5), Inches(1.3))
    center.fill.solid()
    center.fill.fore_color.rgb = COLORS["blue"]
    center.line.fill.background()
    ctf = center.text_frame
    ctf.text = "Schule\n& Wohlbefinden"
    ctf.paragraphs[0].alignment = PP_ALIGN.CENTER
    ctf.paragraphs[0].runs[0].font.size = Pt(16)
    ctf.paragraphs[0].runs[0].font.bold = True
    ctf.paragraphs[0].runs[0].font.name = "Avenir Next"

    card_with_icon(
        s3,
        0.9,
        2.0,
        3.7,
        2.6,
        "🏛️",
        "Strukturmerkmale",
        ["Schulformen", "Notensystem", "Versetzungen", "Abschlüsse"],
        "pink",
    )
    card_with_icon(
        s3,
        8.5,
        1.85,
        3.8,
        2.8,
        "💛",
        "Wohlbefinden",
        ["psychische Gesundheit", "soziale Beziehungen", "Lebenszufriedenheit"],
        "yellow",
    )
    card_with_icon(
        s3,
        4.55,
        4.35,
        4.2,
        2.5,
        "⚖️",
        "Bildungsgerechtigkeit",
        ["faire Chancen", "unabhängig von sozialer Herkunft"],
        "blue",
    )
    set_transition(s3, "cover", {"dir": "r"})

    set_notes(
        s3,
        "Bevor wir in die Analyse einsteigen, klärt die Facharbeit drei Grundbegriffe. "
        "Erstens die Strukturmerkmale von Schule: Dazu gehören unter anderem Schulformen, Noten, Versetzungsregeln und Abschlüsse. "
        "Diese Merkmale organisieren den Schulalltag, wirken aber auch sozial selektiv. "
        "Zweitens das Wohlbefinden: Es umfasst nicht nur Leistungen, sondern auch psychische Gesundheit, soziale Eingebundenheit und Lebenszufriedenheit. "
        "Drittens die Bildungsgerechtigkeit: Gemeint sind faire Bildungszugänge unabhängig von Herkunft. "
        "Gerade im deutschen System zeigt die Literatur eine enge Verbindung zwischen sozialem Hintergrund und Bildungserfolg. "
        "Mit diesen Begriffen im Blick sehen wir uns als Nächstes die Schulformen im deutschen System an.",
    )

    # Slide 4
    s4 = prs.slides.add_slide(blank)
    add_bg(s4, 4)
    text_box(s4, "Das deutsche Schulsystem", 0.8, 0.55, 6.5, 0.8, size=38, bold=True)
    text_box(s4, "Überblick über zentrale Schulformen", 0.8, 1.15, 5.5, 0.4, size=16, color="muted")

    school_cards = [
        ("🎓", "Gymnasium", "Abitur, hohe akademische Anforderungen."),
        ("📘", "Realschule", "Mittlerer Abschluss, Verbindung aus Theorie und Praxis."),
        ("🛠️", "Hauptschule", "Traditionell praxisnah; heute teils in neue Modelle integriert."),
        ("👥", "Gesamtschule", "Längeres gemeinsames Lernen mit verschiedenen Niveaus."),
        ("💼", "Berufliche Schule", "Duale Ausbildung mit enger Verbindung zum Betrieb."),
    ]

    x_positions = [0.8, 3.35, 5.9, 8.45, 11.0]
    for i, (emoji, title, desc) in enumerate(school_cards):
        c = s4.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(x_positions[i]),
            Inches(2.0),
            Inches(2.1),
            Inches(4.7),
        )
        c.fill.solid()
        c.fill.fore_color.rgb = [COLORS["blue"], COLORS["pink"], COLORS["yellow"], COLORS["blue"], COLORS["pink"]][i]
        c.fill.fore_color.brightness = 0.16
        c.line.color.rgb = COLORS["muted"]
        tf = c.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = emoji
        p0.alignment = PP_ALIGN.CENTER
        p0.runs[0].font.size = Pt(28)
        p1 = tf.add_paragraph()
        p1.text = title
        p1.alignment = PP_ALIGN.CENTER
        p1.runs[0].font.size = Pt(14)
        p1.runs[0].font.bold = True
        p1.runs[0].font.name = "Avenir Next"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.runs[0].font.size = Pt(11)
        p2.runs[0].font.name = "Avenir Next"
    set_transition(s4, "uncover", {"dir": "r"})

    set_notes(
        s4,
        "Hier sehen wir die wichtigsten Schulformen, die in der Facharbeit verglichen werden. "
        "Das Gymnasium ist akademisch stark ausgerichtet und führt zum Abitur. "
        "Die Realschule nimmt eine mittlere Position zwischen allgemeinbildender und praxisnäherer Ausrichtung ein. "
        "Die Hauptschule war traditionell stark berufsorientiert, wurde aber in vielen Bundesländern strukturell verändert. "
        "Gesamtschulen verfolgen das Ziel, Schülerinnen und Schüler länger gemeinsam lernen zu lassen. "
        "Berufliche Schulen und das duale System sind ein besonderes Merkmal in Deutschland und international anerkannt. "
        "Diese Vielfalt ist einerseits eine Stärke, andererseits entstehen durch die frühe Aufteilung auch Unterschiede in Chancen und Belastungen. "
        "Damit kommen wir direkt zum Thema Leistungsdruck.",
    )

    # Slide 5
    s5 = prs.slides.add_slide(blank)
    add_bg(s5, 5)
    text_box(s5, "Leistungsdruck und Wohlbefinden", 0.8, 0.55, 7.5, 0.8, size=37, bold=True)
    left_bullets = ["Noten", "Prüfungen", "Konkurrenz", "Erwartungen"]
    right_bullets = ["Stress", "Schlafprobleme", "Unsicherheit", "geringeres Selbstwertgefühl"]
    bullet_box(s5, "Leistungsdruck", left_bullets, 0.9, 1.6, 5.4, 3.9, fill_color="pink")
    bullet_box(s5, "Folgen", right_bullets, 6.9, 1.6, 5.4, 3.9, fill_color="blue")
    text_box(s5, "😵‍💫", 11.5, 0.95, 1.2, 0.8, size=42, align=PP_ALIGN.CENTER)

    memo = s5.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(5.7), Inches(11.2), Inches(1.0))
    memo.fill.solid()
    memo.fill.fore_color.rgb = COLORS["yellow"]
    memo.line.color.rgb = COLORS["muted"]
    mtf = memo.text_frame
    mtf.text = "Merksatz: Leistungsdruck kann Motivation fördern, aber auch das Wohlbefinden beeinträchtigen."
    mtf.paragraphs[0].alignment = PP_ALIGN.CENTER
    mtf.paragraphs[0].runs[0].font.size = Pt(18)
    mtf.paragraphs[0].runs[0].font.bold = True
    mtf.paragraphs[0].runs[0].font.name = "Avenir Next"
    set_transition(s5, "split", {"orient": "vert", "dir": "out"})

    set_notes(
        s5,
        "Die Facharbeit beschreibt Leistungsdruck als einen zentralen Einflussfaktor auf das schulische Wohlbefinden. "
        "Noten und Prüfungen steuern Bildungswege, deshalb werden sie oft als besonders bedeutsam erlebt. "
        "Hinzu kommen Vergleichsprozesse und äußere Erwartungen, etwa von Schule, Elternhaus oder Gesellschaft. "
        "Dabei zeigt sich ein ambivalentes Bild: Leistungsanforderungen können motivierend wirken, wenn sie als machbar erlebt werden. "
        "Werden sie jedoch als dauerhaft überfordernd wahrgenommen, entstehen Stress, Unsicherheit und psychosomatische Belastungen wie Schlafprobleme. "
        "Die Arbeit verweist hier auf WHO- und RKI-Befunde zu schulischem Stress im Jugendalter. "
        "Die entscheidende Frage ist deshalb nicht nur, ob Leistung gefordert wird, sondern wie das schulische Umfeld damit umgeht. "
        "Damit sind wir beim Schulklima.",
    )

    # Slide 6
    s6 = prs.slides.add_slide(blank)
    add_bg(s6, 6)
    text_box(s6, "Schulklima", 0.8, 0.55, 4.0, 0.8, size=40, bold=True)
    text_box(s6, "Beziehungen und Beteiligung als Schutzfaktoren", 0.8, 1.2, 6.0, 0.4, size=16, color="muted")

    card_with_icon(
        s6,
        0.9,
        2.0,
        4.0,
        4.5,
        "👩‍🏫",
        "Lehrer-Schüler-Verhältnis",
        ["Vertrauen senkt Stress", "Unterstützung stärkt Motivation", "Respekt fördert Sicherheit"],
        "blue",
    )
    card_with_icon(
        s6,
        4.95,
        2.0,
        3.9,
        4.5,
        "🗳️",
        "Mitbestimmung",
        ["Beteiligung im Schulalltag", "mehr Selbstwirksamkeit", "Meinung zählt"],
        "pink",
    )
    card_with_icon(
        s6,
        8.9,
        2.0,
        3.9,
        4.5,
        "💬",
        "Feedbackkultur",
        ["Noten plus qualitative Rückmeldung", "individuelle Lernentwicklung", "formative Orientierung"],
        "yellow",
    )
    set_transition(s6, "push", {"dir": "r"})

    set_notes(
        s6,
        "Neben Leistung ist laut Facharbeit vor allem das Schulklima entscheidend. "
        "Ein positives Schulklima umfasst Sicherheit, Zugehörigkeit und verlässliche Beziehungen. "
        "Besonders wichtig ist das Lehrer-Schüler-Verhältnis: Wenn Lernende sich ernst genommen fühlen, steigt die Motivation und Stress kann sinken. "
        "Als zweiter Punkt wird Mitbestimmung betont, zum Beispiel über Klassenrat oder Schülervertretung. "
        "Beteiligung stärkt das Gefühl, wirksam zu sein und nicht nur Objekt von Entscheidungen zu sein. "
        "Drittens spielt Feedback eine wichtige Rolle. Ergänzend zu Noten können qualitative Rückmeldungen Lernfortschritte besser sichtbar machen. "
        "Zusammen zeigen diese Faktoren, dass Wohlbefinden aktiv gestaltbar ist. "
        "Auf der nächsten Folie binden wir das Publikum kurz ein.",
    )

    # Slide 7
    s7 = prs.slides.add_slide(blank)
    add_bg(s7, 7)
    text_box(s7, "Kurze Umfrage", 0.8, 0.55, 4.5, 0.8, size=40, bold=True)
    question = s7.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.55), Inches(11.1), Inches(1.5))
    question.fill.solid()
    question.fill.fore_color.rgb = COLORS["blue"]
    question.fill.fore_color.brightness = 0.1
    question.line.fill.background()
    qtf = question.text_frame
    qtf.text = "Was beeinflusst euer Wohlbefinden in der Schule am stärksten?"
    qtf.paragraphs[0].alignment = PP_ALIGN.CENTER
    qtf.paragraphs[0].runs[0].font.size = Pt(27)
    qtf.paragraphs[0].runs[0].font.bold = True
    qtf.paragraphs[0].runs[0].font.name = "Avenir Next"

    options = [("📚 Notendruck", "pink"), ("👩‍🏫 Lehrkräfte", "yellow"), ("👥 Mitschüler", "blue"), ("🕒 Zeitstress", "pink")]
    pos = [(1.4, 3.45), (6.95, 3.45), (1.4, 4.75), (6.95, 4.75)]
    for (label, c), (x, y) in zip(options, pos):
        b = s7.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(4.9), Inches(1.0))
        b.fill.solid()
        b.fill.fore_color.rgb = COLORS[c]
        b.fill.fore_color.brightness = 0.08
        b.line.color.rgb = COLORS["muted"]
        btf = b.text_frame
        btf.text = label
        btf.paragraphs[0].alignment = PP_ALIGN.CENTER
        btf.paragraphs[0].runs[0].font.size = Pt(21)
        btf.paragraphs[0].runs[0].font.bold = True
        btf.paragraphs[0].runs[0].font.name = "Avenir Next"

    text_box(s7, "Kurzes Handzeichen", 5.2, 6.0, 3.0, 0.4, size=16, color="muted", align=PP_ALIGN.CENTER)
    text_box(
        s7,
        "Alle Faktoren spielen laut Facharbeit eine wichtige Rolle.",
        3.1,
        6.35,
        7.2,
        0.5,
        size=16,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    set_transition(s7, "randomBar", {"dir": "vert"})

    set_notes(
        s7,
        "Hier möchte ich euch kurz aktiv einbinden. "
        "Die Frage lautet: Was beeinflusst euer Wohlbefinden in der Schule am stärksten? "
        "Wir haben vier mögliche Antworten aus den zentralen Themen der Arbeit: Notendruck, Lehrkräfte, Mitschüler und Zeitstress. "
        "Bitte per Handzeichen abstimmen, was ihr persönlich am stärksten wahrnehmt. "
        "Wichtig ist danach die Einordnung: Die Facharbeit zeigt, dass es kein einzelner Faktor ist. "
        "Vielmehr wirken diese Aspekte zusammen und verstärken oder entlasten sich gegenseitig. "
        "Genau deshalb lohnt der Blick über Deutschland hinaus. "
        "Auf der nächsten Folie vergleichen wir internationale Systeme und schauen, welche strukturellen Unterschiede dort sichtbar sind.",
    )

    # Slide 8
    s8 = prs.slides.add_slide(blank)
    add_bg(s8, 8)
    text_box(s8, "Internationaler Vergleich", 0.8, 0.55, 6.2, 0.8, size=38, bold=True)
    text_box(s8, "Deutschland – Finnland – angelsächsische Systeme", 0.8, 1.2, 7.0, 0.4, size=16, color="muted")

    card_with_icon(
        s8,
        0.9,
        1.9,
        3.9,
        4.6,
        "🇩🇪",
        "Deutschland",
        ["frühe Aufteilung", "viele Noten", "starke Selektionslogik"],
        "pink",
    )
    card_with_icon(
        s8,
        4.95,
        1.9,
        3.9,
        4.6,
        "🇫🇮",
        "Finnland",
        ["längeres gemeinsames Lernen", "spätere Differenzierung", "mehr Fokus auf Lernfortschritt"],
        "blue",
    )
    card_with_icon(
        s8,
        9.0,
        1.9,
        3.3,
        4.6,
        "🌍",
        "Angelsächsische Systeme",
        ["mehr Praxis", "mehr Wahlmöglichkeiten", "projektorientierte Elemente"],
        "yellow",
    )
    text_box(s8, "Kein System ist perfekt.", 4.6, 6.55, 4.2, 0.45, size=19, bold=True, align=PP_ALIGN.CENTER)
    set_transition(s8, "circle")

    set_notes(
        s8,
        "Der internationale Vergleich ordnet das deutsche System besser ein. "
        "In Deutschland erfolgt die Aufteilung in Bildungswege meist früh, und Leistungsbewertung durch Noten hat ein hohes Gewicht. "
        "Finnland steht in der Facharbeit für ein Modell mit längerem gemeinsamen Lernen und späterer Differenzierung. "
        "Dort wird in frühen Jahren tendenziell weniger standardisiert getestet und stärker auf individuelle Entwicklung geachtet. "
        "Angelsächsische Systeme zeigen häufig mehr Wahlmöglichkeiten und eine stärkere Praxis- oder Projektorientierung. "
        "Gleichzeitig haben auch diese Systeme Herausforderungen, etwa Unterschiede zwischen Regionen oder Schulen. "
        "Wichtig ist daher die Kernaussage: Es gibt kein perfektes System, aber unterschiedliche strukturelle Prioritäten. "
        "Als Nächstes bündeln wir Stärken und Schwächen des deutschen Modells.",
    )

    # Slide 9
    s9 = prs.slides.add_slide(blank)
    add_bg(s9, 9)
    text_box(s9, "Stärken und Schwächen", 0.8, 0.55, 6.0, 0.8, size=39, bold=True)
    bullet_box(
        s9,
        "Stärken",
        ["verschiedene Bildungswege", "duales Ausbildungssystem", "hohe akademische Standards"],
        0.9,
        1.8,
        5.7,
        4.8,
        fill_color="green_soft",
    )
    bullet_box(
        s9,
        "Schwächen",
        ["frühe Selektion", "Leistungsdruck", "soziale Ungleichheiten"],
        6.75,
        1.8,
        5.7,
        4.8,
        fill_color="pink",
    )
    set_transition(s9, "wipe", {"dir": "r"})

    set_notes(
        s9,
        "Die Facharbeit bewertet das deutsche Schulsystem differenziert. "
        "Zu den Stärken zählt die Vielfalt an Bildungswegen, durch die unterschiedliche Lerninteressen und Leistungsprofile adressiert werden können. "
        "Besonders hervorgehoben wird das duale Ausbildungssystem als gute Brücke in den Arbeitsmarkt. "
        "Auch die akademischen Standards, vor allem im gymnasialen Bereich, gelten als qualitativ anspruchsvoll. "
        "Gleichzeitig stehen diesen Vorteilen klare Schwächen gegenüber: die frühe Selektion, ein oft hoher Leistungsdruck und die enge Kopplung zwischen sozialer Herkunft und Bildungserfolg. "
        "Diese Schwächen sind für das Wohlbefinden besonders relevant, weil sie Stress und Ungleichheitswahrnehmung verstärken können. "
        "Daraus leiten sich Reformideen ab, die wir nun anschauen.",
    )

    # Slide 10
    s10 = prs.slides.add_slide(blank)
    add_bg(s10, 10)
    text_box(s10, "Reformansätze", 0.8, 0.55, 4.5, 0.8, size=40, bold=True)
    text_box(s10, "Strukturen in Richtung Zukunft weiterentwickeln", 0.8, 1.2, 7.5, 0.45, size=16, color="muted")
    text_box(s10, "➡️", 11.4, 0.45, 1.2, 0.8, size=48, align=PP_ALIGN.CENTER)

    circles = [
        ("🧭", "Mehr Lebenskompetenzen", "z. B. Finanzbildung,\nMedienkompetenz,\nAlltagsbezug"),
        ("🕰️", "Spätere Leistungs-\ndifferenzierung", "mehr Zeit zur\nindividuellen Entwicklung\nvor der Aufteilung"),
        ("📝", "Alternative\nBewertungsformen", "mehr Lernfeedback,\nPortfolio- und\nprojektbezogene Bewertung"),
    ]
    cx = [1.2, 4.9, 8.6]
    fills = ["blue", "pink", "yellow"]
    for i, (emoji, title, desc) in enumerate(circles):
        c = s10.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(cx[i]), Inches(2.0), Inches(3.3), Inches(3.3))
        c.fill.solid()
        c.fill.fore_color.rgb = COLORS[fills[i]]
        c.fill.fore_color.brightness = 0.1
        c.line.color.rgb = COLORS["muted"]
        tf = c.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = emoji
        p0.alignment = PP_ALIGN.CENTER
        p0.runs[0].font.size = Pt(26)
        p1 = tf.add_paragraph()
        p1.text = title
        p1.alignment = PP_ALIGN.CENTER
        p1.runs[0].font.size = Pt(14)
        p1.runs[0].font.bold = True
        p1.runs[0].font.name = "Avenir Next"
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.alignment = PP_ALIGN.CENTER
        p2.runs[0].font.size = Pt(11)
        p2.runs[0].font.name = "Avenir Next"
    set_transition(s10, "wheel", {"spokes": "1"})

    set_notes(
        s10,
        "Aus den analysierten Schwächen leitet die Facharbeit drei zentrale Reformansätze ab. "
        "Erstens: mehr Lebenskompetenzen im Curriculum, etwa Finanzbildung oder Medienkompetenz. "
        "Damit soll Schule alltagsnäher werden und junge Menschen besser auf selbstständiges Handeln vorbereiten. "
        "Zweitens: spätere Leistungsdifferenzierung, damit Bildungsentscheidungen weniger früh und auf einer breiteren Entwicklungsbasis getroffen werden. "
        "Drittens: alternative Bewertungsformen, die Lernprozesse stärker begleiten, statt nur punktuelle Prüfungsleistungen zu bewerten. "
        "Diese Ansätze verstehen Leistung nicht als Gegensatz zum Wohlbefinden, sondern als etwas, das durch gute Strukturen nachhaltiger gefördert werden kann. "
        "Damit kommen wir zum abschließenden Fazit.",
    )

    # Slide 11
    s11 = prs.slides.add_slide(blank)
    add_bg(s11, 11)
    text_box(s11, "Schule beeinflusst mehr als nur Noten.", 0.8, 0.7, 9.4, 0.9, size=42, bold=True)
    summary = s11.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(2.0), Inches(8.8), Inches(3.6))
    summary.fill.solid()
    summary.fill.fore_color.rgb = COLORS["blue"]
    summary.fill.fore_color.brightness = 0.2
    summary.line.fill.background()
    stf = summary.text_frame
    stf.clear()
    statements = [
        "✓ Strukturen wirken auf das Wohlbefinden.",
        "✓ Leistungsdruck und Schulklima spielen eine zentrale Rolle.",
        "✓ Reformen könnten Bildung und Glück stärker miteinander verbinden.",
    ]
    p = stf.paragraphs[0]
    p.text = statements[0]
    p.runs[0].font.size = Pt(23)
    p.runs[0].font.bold = True
    p.runs[0].font.name = "Avenir Next"
    p.runs[0].font.color.rgb = COLORS["dark"]
    for s in statements[1:]:
        px = stf.add_paragraph()
        px.text = s
        px.runs[0].font.size = Pt(21)
        px.runs[0].font.bold = True
        px.runs[0].font.name = "Avenir Next"
        px.runs[0].font.color.rgb = COLORS["dark"]
        px.space_after = Pt(8)

    calm = s11.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(9.95), Inches(2.0), Inches(2.7), Inches(3.6))
    calm.fill.solid()
    calm.fill.fore_color.rgb = COLORS["pink"]
    calm.fill.fore_color.brightness = 0.22
    calm.line.fill.background()
    text_box(s11, "🌅", 10.5, 2.55, 1.8, 0.9, size=52, align=PP_ALIGN.CENTER)
    text_box(s11, "ruhiger Ausblick", 10.2, 3.5, 2.2, 0.4, size=14, color="muted", align=PP_ALIGN.CENTER)

    text_box(s11, "Danke für eure Aufmerksamkeit!", 0.95, 6.2, 7.2, 0.6, size=28, bold=True)
    set_transition(s11, "fade")

    set_notes(
        s11,
        "Zusammenfassend zeigt die Facharbeit klar: Schule beeinflusst deutlich mehr als Noten. "
        "Strukturelle Merkmale wie Schulform, Selektionszeitpunkt und Bewertungssystem wirken auf Lernbedingungen und Wohlbefinden zugleich. "
        "Zudem wurde sichtbar, dass Leistungsdruck und Schulklima zentrale Stellschrauben sind, wenn es um Motivation, Selbstbild und psychische Belastung geht. "
        "Der internationale Vergleich macht deutlich, dass andere Wege möglich sind, auch wenn kein System perfekt ist. "
        "Daraus ergeben sich Reformoptionen, die Bildungsgerechtigkeit und Wohlbefinden besser verbinden könnten. "
        "Für die Leitfrage bedeutet das: Ja, strukturelle Unterschiede im deutschen Schulsystem beeinflussen das Wohlbefinden von Schülerinnen und Schülern spürbar. "
        "Vielen Dank für eure Aufmerksamkeit, ich freue mich auf eure Fragen und eure Eindrücke.",
    )

    output_file = "Seminarfach_Glueck_Keynote_Style.pptx"
    prs.save(output_file)
    print(f"Created: {output_file}")


if __name__ == "__main__":
    build()
